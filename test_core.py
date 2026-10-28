"""核心逻辑自动化测试（开发用）。

用法：python test_core.py
"""

import shutil
import sys
from pathlib import Path

import excel_core

BASE = Path(__file__).parent
SAMPLES = BASE / 'test_samples'
OUT = BASE / 'test_out'


def log(msg):
    print(msg)


def check(cond, name):
    print(('PASS' if cond else 'FAIL') + ' | ' + name)
    if not cond:
        sys.exit(1)


def clean_out():
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir()


def test_rename():
    clean_out()
    files = [str(SAMPLES / f) for f in ('数据1.xlsx', '数据2.xlsx', '数据3.xlsx',
                                        '数据4.xlsx', '数据5.xlsx', '数据6.xlsx', '数据7.xlsx')]
    rule = [('cell', 'A1'), ('cell', 'B2'), ('cell', 'C3')]
    plan, errors = excel_core.plan_rename(files, rule, '-', OUT)
    print('--- 重命名计划 ---')
    names = [p.name for p in plan]
    for p in plan:
        print(f'  {p.src.name} -> {p.name}')
    check(not errors, '无读取错误')

    expect = {
        '数据1': '张三-20240115-98.5.xlsx',
        '数据2': '李四-20240220-100.xlsx',
        '数据3': '李四-20240220-100 (1).xlsx',      # 重名去重
        '数据4': '王五_Q_A-20240301-7.xlsx',        # 非法字符替换
        '数据6': '赵六-7.xlsx',                     # 空值跳过
        # 数据7：'CON-20240501-9' 整体并非保留名，Windows 允许，故不加前缀
        '数据7': 'CON-20240501-9.xlsx',
    }
    for p in plan:
        stem = p.src.stem
        if stem in expect:
            check(p.name == expect[stem], f'{stem} 新文件名 {p.name} == {expect[stem]}')
    # 数据5 超长截断到 150 字符；macOS 还要满足 APFS 的 255 字节文件名上限
    d5 = next(p for p in plan if p.src.stem == '数据5')
    expect_d5_stem = '很' * 150
    if sys.platform == 'darwin':
        expect_d5_stem = expect_d5_stem.encode('utf-8')[:240].decode('utf-8', errors='ignore')
    check(d5.name == expect_d5_stem + '.xlsx',
          f'超长截断（实际 {len(Path(d5.name).stem)} 字符）')

    # 纯保留设备名（单单元格）应加下划线前缀
    check(excel_core.sanitize_name_part('CON') == '_CON', '保留名 CON 前缀下划线')
    check(excel_core.sanitize_name_part('CON.txt') == '_CON.txt', '保留名带扩展名前缀下划线')

    # 执行
    summary = excel_core.execute_rename(plan, log)
    print('--- 执行结果 ---', summary)
    check(summary['ok'] == 7, '7 个文件全部成功')
    check(all(p.dst.exists() for p in plan), '所有输出文件已生成')
    check(not (SAMPLES / '数据1.xlsx' in [p.dst for p in plan]), '原文件未被覆盖')

    # 二次执行应触发磁盘重名检测（已存在的基础名/序号名会被递增跳过）
    plan2, _ = excel_core.plan_rename(files, rule, '-', OUT)
    names2 = {p.name for p in plan2}
    check(len(names2) == len(plan2), '二次计划文件名全部唯一')
    check(all(' (' in p.name for p in plan2
              if p.src.stem in ('数据1', '数据2', '数据3', '数据4', '数据6', '数据7')),
          '磁盘已存在时自动追加序号')


def test_rename_advanced():
    """自由组合规则：手打文字 / 单元格 / 原文件名 混排。"""
    out = BASE / 'test_out_adv'
    if out.exists():
        shutil.rmtree(out)
    out.mkdir()

    # 1) 文字 + 单元格 + 原文件名 混排，分隔符插在相邻 data 之间
    files = [str(SAMPLES / '数据1.xlsx')]   # A1=张三, B2=2024-01-15
    rule = [('text', '客户_'), ('cell', 'A1'), ('filename', ''), ('text', '_已审')]
    plan, errors = excel_core.plan_rename(files, rule, '-', out)
    check(not errors, '高级规则无读取错误')
    check(plan[0].name == '客户_张三-数据1_已审.xlsx', f'混排规则 {plan[0].name}')

    # 2) 文字与单元格相邻时，不插分隔符；单元格之间才插
    rule2 = [('text', '客户'), ('cell', 'A1'), ('cell', 'B2')]
    plan, _ = excel_core.plan_rename(files, rule2, '-', out)
    check(plan[0].name == '客户张三-20240115.xlsx', f'文字相邻不插分隔符 {plan[0].name}')

    # 3) 单元格全空时回退原文件名，文字仍保留
    empty = out / '空白.xlsx'
    from openpyxl import Workbook
    wb = Workbook()
    wb.save(empty)
    rule3 = [('text', '前缀'), ('cell', 'A1'), ('text', '后缀')]
    plan, _ = excel_core.plan_rename([str(empty)], rule3, '-', out)
    check(plan[0].name == '前缀后缀.xlsx', f'空单元格回退原文件名+文字 {plan[0].name}')


def test_fill():
    clean_out()
    template = str(SAMPLES / '合同模板.xlsx')
    files = [str(SAMPLES / f) for f in ('数据1.xlsx', '数据2.xlsx')]
    mapping = [('cell', 'A1', 'B2'), ('cell', 'B2', 'C5'), ('cell', 'C3', 'D6')]
    plan, errors = excel_core.plan_fill(files, template, mapping, OUT)
    print('--- 填充计划 ---')
    for p in plan:
        print(f'  {p.src.name} -> {p.dst.name}')
    check(not errors, '无读取错误')
    check(plan[0].dst.name == '合同模板_数据1.xlsx', '输出文件名 = 模板名 + 源文件名')

    summary = excel_core.execute_fill(template, plan, mapping, log)
    print('--- 填充结果 ---', summary)
    check(summary['ok'] == 2, '2 个文件全部成功')
    check(all(p.dst.exists() for p in plan), '所有输出文件已生成')

    # 验证内容与格式
    from openpyxl import load_workbook
    wb = load_workbook(plan[0].dst)
    ws = wb.worksheets[0]
    check(ws['B2'].value == '张三', '模板 B2 填入源 A1（张三）')
    from datetime import datetime
    check(ws['C5'].value == datetime(2024, 1, 15), '模板 C5 填入日期（保持日期类型）')
    check(ws['D6'].value == 98.5, '模板 D6 填入数值（保持数值类型）')
    check(ws['B2'].number_format == '0.00', '模板 B2 数字格式保留')
    check(ws['C5'].number_format == 'yyyy-mm-dd', '模板 C5 日期格式保留')
    check(ws['A1'].font.bold, '模板标题加粗保留')
    check(ws['A1'].value == '合同信息表', '模板其他内容未受影响')
    wb.close()

    # 映射验证：数据2 的日期写入 C5
    wb2 = load_workbook(plan[1].dst)
    ws2 = wb2.worksheets[0]
    check(ws2['C5'].value == datetime(2024, 2, 20), '数据2 日期正确')
    wb2.close()


def test_fill_column():
    """整列映射：源列数据填入模板列；目标已有内容时插入空白行下移，不覆盖。"""
    from openpyxl import Workbook, load_workbook
    clean_out()

    # 源文件：B 列 2..6 行 = 产品A..产品E，第 7 行为空（列到此为止）
    src = OUT / '列源.xlsx'
    wb = Workbook()
    ws = wb.active
    for r, v in enumerate(['产品A', '产品B', '产品C', '产品D', '产品E'], start=2):
        ws[f'B{r}'] = v
    wb.save(src)

    # 模板：C 列 5..7 行已有内容
    tpl = OUT / '列模板.xlsx'
    wb = Workbook()
    ws = wb.active
    for r, v in zip((5, 6, 7), ('已有1', '已有2', '已有3')):
        ws[f'C{r}'] = v
    wb.save(tpl)

    mapping = [('column', 'B', 2, 'C', 5)]
    plan, errors = excel_core.plan_fill([str(src)], str(tpl), mapping, OUT)
    check(not errors, '整列映射无读取错误')
    summary = excel_core.execute_fill(str(tpl), plan, mapping, log)
    check(summary['ok'] == 1, '整列填充成功')

    out_wb = load_workbook(plan[0].dst)
    ws = out_wb.worksheets[0]
    for r, v in enumerate(['产品A', '产品B', '产品C', '产品D', '产品E'], start=5):
        check(ws[f'C{r}'].value == v, f'C{r} 填入 {v}')
    check(ws['C10'].value == '已有1', '原有内容1 下移不覆盖')
    check(ws['C11'].value == '已有2', '原有内容2 下移不覆盖')
    check(ws['C12'].value == '已有3', '原有内容3 下移不覆盖')
    out_wb.close()

    # 无冲突情况：模板列为空，直接填充，不插入多余行
    tpl2 = OUT / '列模板2.xlsx'
    wb = Workbook()
    wb.save(tpl2)
    plan2, _ = excel_core.plan_fill([str(src)], str(tpl2), mapping, OUT)
    excel_core.execute_fill(str(tpl2), plan2, mapping, log)
    ws2 = load_workbook(plan2[0].dst).worksheets[0]
    check(ws2['C5'].value == '产品A' and ws2['C9'].value == '产品E', '无冲突时直接填充')
    check(ws2['C10'].value is None, '无冲突时不插入多余行')


def test_fill_column_row_boundary():
    """列填充扩展：源数据超出空白时在目标列边界处插入，且不影响其他列数据对齐。"""
    from openpyxl import Workbook as WX, load_workbook as LW
    clean_out()

    # 源：B 列 5..9 = 产品A..E
    src = OUT / 's.xlsx'
    wb = WX()
    ws = wb.active
    for r, v in enumerate(['产品A', '产品B', '产品C', '产品D', '产品E'], start=5):
        ws[f'B{r}'] = v
    wb.save(src)

    # --- .xlsx：填 B 列起始行2；B2-B4 空白，B5=100(目标列边界)；A 列有序号 1,2,3 和合计 ---
    tpl = OUT / 't.xlsx'
    wb = WX()
    ws = wb.active
    ws['A1'] = '序号'
    ws['B1'] = '金额'
    ws['A2'] = 1
    ws['A3'] = 2
    ws['A4'] = 3
    ws['A5'] = '合计'
    ws['B5'] = 100
    wb.save(tpl)
    mapping = [('column', 'B', 5, 'B', 2)]
    plan, _ = excel_core.plan_fill([str(src)], str(tpl), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl), plan, mapping, log)
    check(summary['ok'] == 1, 'xlsx 列扩展填充成功')
    ws2 = LW(plan[0].dst).worksheets[0]
    check(ws2['B2'].value == '产品A' and ws2['B6'].value == '产品E', 'xlsx B2-B6 填满源列')
    check(ws2['B7'].value == 100, 'xlsx 边界值推下到 B7')
    check(ws2['A7'].value == '合计', 'xlsx 合计随边界下移')
    check(ws2['A2'].value == 1 and ws2['A3'].value == 2 and ws2['A4'].value == 3,
          'xlsx 其他列(序号)保持对齐不被推下')

    # --- .xls：同样场景 ---
    import xlwt
    tpl_x = OUT / 't.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('S')
    ws.write(0, 0, '序号')
    ws.write(0, 1, '金额')
    ws.write(1, 0, 1)
    ws.write(2, 0, 2)
    ws.write(3, 0, 3)
    ws.write(4, 0, '合计')
    ws.write(4, 1, 100)   # 0基第4行：B5=100，目标列边界
    wb.save(str(tpl_x))
    plan, _ = excel_core.plan_fill([str(src)], str(tpl_x), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl_x), plan, mapping, log)
    check(summary['ok'] == 1, 'xls 列扩展填充成功')
    import xlrd
    s = xlrd.open_workbook(str(plan[0].dst)).sheet_by_index(0)
    check(s.cell_value(1, 1) == '产品A' and s.cell_value(5, 1) == '产品E', 'xls B2-B6 填满源列')
    check(s.cell_value(6, 1) == 100, 'xls 边界值推下到第7行(0基6)')
    check(s.cell_value(6, 0) == '合计', 'xls 合计随边界下移')
    check(s.cell_value(1, 0) == 1 and s.cell_value(2, 0) == 2 and s.cell_value(3, 0) == 3,
          'xls 其他列(序号)保持对齐')


def test_fill_column_merged():
    """列填充不得覆盖模板原有数据或合并单元格：合并区视为边界，随插入整体下移。"""
    from openpyxl import Workbook as WX, load_workbook as LW
    clean_out()

    # 源：B 列 5..8 = 产品A..D
    src = OUT / 's.xlsx'
    wb = WX()
    ws = wb.active
    for r, v in enumerate(['产品A', '产品B', '产品C', '产品D'], start=5):
        ws[f'B{r}'] = v
    wb.save(src)

    # --- .xlsx：目标列 B2:B3 合并('原有合并')，B6='边界内容' ---
    tpl = OUT / 't.xlsx'
    wb = WX()
    ws = wb.active
    ws.merge_cells('B2:B3')
    ws['B2'] = '原有合并'
    ws['B6'] = '边界内容'
    wb.save(tpl)
    mapping = [('column', 'B', 5, 'B', 2)]
    plan, _ = excel_core.plan_fill([str(src)], str(tpl), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl), plan, mapping, log)
    check(summary['ok'] == 1, 'xlsx 合并单元格列填充成功')
    ws2 = LW(plan[0].dst).worksheets[0]
    check(ws2['B2'].value == '产品A' and ws2['B5'].value == '产品D', 'xlsx 源列填满不丢失')
    check(ws2['B6'].value == '原有合并', 'xlsx 合并内容未被覆盖')
    ranges = [str(r) for r in ws2.merged_cells.ranges]
    check('B6:B7' in ranges, f'xlsx 合并区域随插入下移 ({ranges})')
    check(ws2['B10'].value == '边界内容', 'xlsx 边界内容下移不覆盖')

    # --- .xls：同样场景 ---
    import xlwt
    tpl_x = OUT / 't.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('S')
    ws.write_merge(1, 2, 1, 1, '原有合并')   # B2:B3
    ws.write(5, 1, '边界内容')                # B6
    wb.save(str(tpl_x))
    plan, _ = excel_core.plan_fill([str(src)], str(tpl_x), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl_x), plan, mapping, log)
    check(summary['ok'] == 1, 'xls 合并单元格列填充成功')
    import xlrd
    s = xlrd.open_workbook(str(plan[0].dst), formatting_info=True).sheet_by_index(0)
    check(s.cell_value(1, 1) == '产品A' and s.cell_value(4, 1) == '产品D', 'xls 源列填满不丢失')
    check(s.cell_value(5, 1) == '原有合并', 'xls 合并内容未被覆盖')
    check((5, 7, 1, 2) in [tuple(m) for m in s.merged_cells], 'xls 合并区域随插入下移')
    check(s.cell_value(9, 1) == '边界内容', 'xls 边界内容下移不覆盖')


def test_xls():
    """旧版 .xls 支持：重命名读取、.xls 源填充、.xls 模板重建填充（含插入行）。"""
    import xlrd
    import xlwt
    from datetime import datetime
    from openpyxl import Workbook, load_workbook
    clean_out()

    # --- 创建 .xls 源：A1=张三, B2=日期, C3=98.5；B 列 5..7 行=产品A..C ---
    src_xls = OUT / '旧数据.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sheet1')
    ws.write(0, 0, '张三')
    date_style = xlwt.easyxf('', num_format_str='YYYY-MM-DD')
    ws.write(1, 1, datetime(2024, 1, 15), date_style)
    ws.write(2, 2, 98.5)
    for r, v in enumerate(['产品A', '产品B', '产品C'], start=4):
        ws.write(r, 1, v)
    wb.save(str(src_xls))

    # --- 1) .xls 重命名 ---
    plan, errors = excel_core.plan_rename([str(src_xls)],
                                          [('cell', 'A1'), ('cell', 'B2'), ('cell', 'C3')], '-', OUT)
    check(not errors, 'xls 重命名无读取错误')
    check(plan[0].name == '张三-20240115-98.5.xls', f'xls 重命名 {plan[0].name}')

    # 规则引用超出范围的单元格：视为空，不应报 list index out of range
    plan, errors = excel_core.plan_rename([str(src_xls)], [('cell', 'A1'), ('cell', 'Z99')], '-', OUT)
    check(not errors, 'xls 越界读取不报错')
    check(plan[0].name == '张三.xls', f'xls 越界视为空 {plan[0].name}')

    # .xls 源整列读取（B 列 5..7 行），数据读到工作表边界不报错
    mapping4 = [('column', 'B', 5, 'D', 2)]
    tpl4 = str(SAMPLES / '合同模板.xlsx')
    plan4, _ = excel_core.plan_fill([str(src_xls)], tpl4, mapping4, OUT)
    summary4 = excel_core.execute_fill(tpl4, plan4, mapping4, log)
    check(summary4['ok'] == 1, 'xls 源整列填充成功')
    ws4 = load_workbook(plan4[0].dst).worksheets[0]
    check(ws4['D2'].value == '产品A' and ws4['D4'].value == '产品C', 'xls 源整列值正确')

    # --- 2) .xls 源 → .xlsx 模板（单元格映射） ---
    tpl = str(SAMPLES / '合同模板.xlsx')
    mapping = [('cell', 'A1', 'B2'), ('cell', 'B2', 'C5'), ('cell', 'C3', 'D6')]
    plan, _ = excel_core.plan_fill([str(src_xls)], tpl, mapping, OUT)
    summary = excel_core.execute_fill(tpl, plan, mapping, log)
    check(summary['ok'] == 1, 'xls 源→xlsx 模板填充成功')
    ws = load_workbook(plan[0].dst).worksheets[0]
    check(ws['B2'].value == '张三', 'xls 源 A1 填入模板 B2')
    check(ws['C5'].value == datetime(2024, 1, 15), 'xls 源 B2 日期填入模板 C5')
    check(ws['D6'].value == 98.5, 'xls 源 C3 数值填入模板 D6')

    # --- 3) .xls 源 → .xls 模板（单元格映射，格式保留） ---
    tpl_xls = OUT / '旧模板.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sheet1')
    border = xlwt.easyxf('borders: left thin, right thin, top thin, bottom thin')
    ws.write(1, 1, None, border)                       # B2 带边框
    ws.write(4, 2, None, date_style)                   # C5 日期格式
    wb.save(str(tpl_xls))
    mapping2 = [('cell', 'A1', 'B2'), ('cell', 'B2', 'C5')]
    plan2, _ = excel_core.plan_fill([str(src_xls)], str(tpl_xls), mapping2, OUT)
    summary2 = excel_core.execute_fill(str(tpl_xls), plan2, mapping2, log)
    check(summary2['ok'] == 1, 'xls→xls 模板填充成功')
    rb = xlrd.open_workbook(str(plan2[0].dst), formatting_info=True)
    s = rb.sheet_by_index(0)
    check(s.cell_value(1, 1) == '张三', 'xls 模板 B2 填入张三')
    xf_b2 = rb.xf_list[s.cell_xf_index(1, 1)]
    check(xf_b2.border.left_line_style == 1, 'xls 模板 B2 边框保留')
    date_val = xlrd.xldate.xldate_as_datetime(s.cell_value(4, 2), rb.datemode)
    check(date_val == datetime(2024, 1, 15), 'xls 模板 C5 日期保留')
    rb.release_resources()

    # --- 4) .xls 模板 整列填充（插入空白行下移，不覆盖） ---
    tpl_xls2 = OUT / '旧表格模板.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sheet1')
    for r, v in zip((4, 5, 6), ('产品', '旧项1', '旧项2')):
        ws.write(r, 0, v, border)   # A5/A6/A7
    wb.save(str(tpl_xls2))
    mapping3 = [('column', 'B', 2, 'A', 6)]
    plan3, _ = excel_core.plan_fill([str(SAMPLES / '列数据源.xlsx')], str(tpl_xls2), mapping3, OUT)
    summary3 = excel_core.execute_fill(str(tpl_xls2), plan3, mapping3, log)
    check(summary3['ok'] == 1, 'xls 模板整列填充成功')
    s3 = xlrd.open_workbook(str(plan3[0].dst)).sheet_by_index(0)
    check(s3.cell_value(5, 0) == '产品A', 'xls 模板 A6 填入产品A')
    check(s3.cell_value(9, 0) == '产品E', 'xls 模板 A10 填入产品E')
    check(s3.cell_value(10, 0) == '旧项1', 'xls 模板 旧项1 下移不覆盖')
    check(s3.cell_value(11, 0) == '旧项2', 'xls 模板 旧项2 下移不覆盖')

    # --- 5) .xls 模板含合并单元格标题：填充不应报 overwrite cell 错 ---
    tpl_m = OUT / '对账模板.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('对账单模板')
    header = xlwt.easyxf('font: bold on; borders: left thin, right thin, top thin, bottom thin; align: horiz center')
    ws.write_merge(0, 0, 0, 3, '对账单', header)   # A1:D1 合并标题
    ws.write(2, 1, None, border)                   # B3 待填
    wb.save(str(tpl_m))
    mapping_m = [('cell', 'A1', 'B3')]
    plan_m, _ = excel_core.plan_fill([str(SAMPLES / '数据1.xlsx')], str(tpl_m), mapping_m, OUT)
    summary_m = excel_core.execute_fill(str(tpl_m), plan_m, mapping_m, log)
    check(summary_m['ok'] == 1, 'xls 合并标题模板填充成功（无 overwrite cell 错）')
    rb_m = xlrd.open_workbook(str(plan_m[0].dst), formatting_info=True)
    s_m = rb_m.sheet_by_index(0)
    check(s_m.cell_value(0, 0) == '对账单', '合并标题保留')
    check((0, 1, 0, 4) in [tuple(m) for m in s_m.merged_cells], '合并单元格保留')
    check(s_m.cell_value(2, 1) == '张三', 'B3 填入张三')
    rb_m.release_resources()

    # --- 6) 内置数字格式：xlrd 对部分内置格式返回 format_str=None，转换不得产生 None ---
    rb6 = xlrd.open_workbook(str(src_xls), formatting_info=True)
    for fk, fmt in rb6.format_map.items():
        check(excel_core._resolve_num_format(rb6, fk) is not None,
              f'format_key {fk} 解析非 None')
    for fk in (27, 28, 29, 30, 50, 51):  # xlrd 返回 None format_str 的内置 key
        check(excel_core._resolve_num_format(rb6, fk) is not None,
              f'内置格式 {fk} 回退非 None')
    # 模拟 format_str=None 时，_make_xlwt_style 也不得产生 None num_format_str
    class _Fmt:
        format_str = None
    class _Xf:
        font_index = 0
        format_key = 27
        border = type('B', (), {'left_line_style': 1, 'right_line_style': 1,
                                'top_line_style': 1, 'bottom_line_style': 1})()
        background = type('Bg', (), {'fill_pattern': 0, 'pattern_colour_index': 0,
                                     'background_colour_index': 0})()
        alignment = type('Al', (), {'hor_align': 0, 'vert_align': 0, 'text_wrapped': 0})()
    class _Font:
        name = 'Arial'
        bold = 0
        italic = 0
        height = 200
        colour_index = 32767
    class _Book:
        xf_list = [_Xf()]
        font_list = [_Font()]
        format_map = {27: _Fmt()}
    st = excel_core._make_xlwt_style(_Book(), 0)
    check(st.num_format_str is not None, '防护后 num_format_str 非 None')
    check(st.num_format_str == 'General', 'None 格式串回退为 General')
    rb6.release_resources()

    # --- 7) 带货币内置格式的 .xls 模板填充不报错 ---
    tpl_cur = OUT / '货币模板.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('对账单模板')
    ws.write(2, 1, None, xlwt.easyxf('', num_format_str='#,##0.00'))
    wb.save(str(tpl_cur))
    plan_cur, _ = excel_core.plan_fill([str(src_xls)], str(tpl_cur), [('cell', 'C3', 'B3')], OUT)
    sum_cur = excel_core.execute_fill(str(tpl_cur), plan_cur, [('cell', 'C3', 'B3')], log)
    check(sum_cur['ok'] == 1, '货币格式模板填充成功')


def test_fill_merged_cell():
    """映射到合并单元格：目标在合并区非左上角时自动重定向到锚点，值可见不丢失。"""
    from openpyxl import Workbook as WX, load_workbook as LW
    clean_out()

    # 源
    src = OUT / 'src.xlsx'
    wb = WX()
    wb.active['A1'] = '张三'
    wb.save(src)

    # --- .xlsx：B2:C2 合并，映射源 A1 → C2（非锚点） ---
    tpl = OUT / 'm.xlsx'
    wb = WX()
    ws = wb.active
    ws.merge_cells('B2:C2')
    ws['B2'] = '原值'
    wb.save(tpl)
    mapping = [('cell', 'A1', 'C2')]
    plan, _ = excel_core.plan_fill([str(src)], str(tpl), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl), plan, mapping, log)
    check(summary['ok'] == 1, 'xlsx 映射到合并区非锚点不再报错')
    ws2 = LW(plan[0].dst).worksheets[0]
    check(ws2['B2'].value == '张三', 'xlsx 值写入合并区锚点 B2')

    # --- .xls：B2:C2 合并，映射源 A1 → C2（非锚点） ---
    import xlwt
    tpl_x = OUT / 'm.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('S')
    ws.write_merge(1, 1, 1, 2, '原值')   # B2:C2
    wb.save(str(tpl_x))
    plan, _ = excel_core.plan_fill([str(src)], str(tpl_x), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl_x), plan, mapping, log)
    check(summary['ok'] == 1, 'xls 映射到合并区非锚点成功')
    import xlrd
    s = xlrd.open_workbook(str(plan[0].dst), formatting_info=True).sheet_by_index(0)
    check(s.cell_value(1, 1) == '张三', 'xls 值写入合并区锚点 B2(0基1,1)')
    # 合并区非锚点 C2 不应有独立值（可能未计入使用范围）
    if s.ncols > 2:
        check(s.cell_value(1, 2) == '', 'xls 合并区非锚点 C2 无独立值')
    check((1, 2, 1, 3) in [tuple(m) for m in s.merged_cells], 'xls 合并单元格保留')


def test_fill_sum():
    """列求和：数值/数值型字符串/前后空格/千分位 → 准确转换求和，填入指定或合并单元格。"""
    from openpyxl import Workbook as WX, load_workbook as LW
    clean_out()

    # 数值转换边界
    check(excel_core._to_number(' 100 ') == 100.0, '字符型数字去空格 100')
    check(excel_core._to_number('1,500.5') == 1500.5, '千分位 1,500.5')
    check(excel_core._to_number(5) == 5.0, '数字 5')
    check(excel_core._to_number('abc') is None, '非数字文本 → None')
    check(excel_core._to_number('') is None, '空串 → None')
    check(excel_core._to_number(True) is None, '布尔 → None')

    # 源 .xlsx：B 列 2..6 行 = 100, ' 200 ', 300, '1,500.5', '作废'（非数值跳过），第7行空
    src = OUT / 'sum.xlsx'
    wb = WX()
    ws = wb.active
    ws['B2'] = 100
    ws['B3'] = ' 200 '
    ws['B4'] = 300
    ws['B5'] = '1,500.5'
    ws['B6'] = '作废'
    wb.save(src)

    total, count = excel_core.sum_column(str(src), 'B', 2)
    check(abs(total - 2100.5) < 1e-9, f'求和 = 2100.5（实际 {total}）')
    check(count == 4, f'有效数值 4 个（实际 {count}）')

    # 求和填入模板普通单元格
    tpl = OUT / 't.xlsx'
    wb = WX()
    ws = wb.active
    ws['D5'] = None
    ws['D5'].number_format = '#,##0.00'
    wb.save(tpl)
    mapping = [('sum', 'B', 2, 'D5')]
    plan, _ = excel_core.plan_fill([str(src)], str(tpl), mapping, OUT)
    summary = excel_core.execute_fill(str(tpl), plan, mapping, log)
    check(summary['ok'] == 1, '求和填充成功')
    ws2 = LW(plan[0].dst).worksheets[0]
    check(abs(ws2['D5'].value - 2100.5) < 1e-9, f'D5 = 2100.5（实际 {ws2["D5"].value}）')
    check(ws2['D5'].number_format == '#,##0.00', '模板数字格式保留')

    # 求和填入合并单元格（非锚点自动定位到锚点）
    tpl2 = OUT / 't2.xlsx'
    wb = WX()
    ws = wb.active
    ws.merge_cells('D5:E5')
    ws['D5'] = None
    wb.save(tpl2)
    plan, _ = excel_core.plan_fill([str(src)], str(tpl2), [('sum', 'B', 2, 'E5')], OUT)
    excel_core.execute_fill(str(tpl2), plan, [('sum', 'B', 2, 'E5')], log)
    ws2 = LW(plan[0].dst).worksheets[0]
    check(abs(ws2['D5'].value - 2100.5) < 1e-9, f'合并单元格锚点 D5 = 2100.5（实际 {ws2["D5"].value}）')

    # .xls 源
    import xlwt
    src_x = OUT / 'sum.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('S')
    for r, v in zip((1, 2, 3), (' 100 ', 200, '300.5')):
        ws.write(r, 1, v)
    wb.save(str(src_x))
    total, count = excel_core.sum_column(str(src_x), 'B', 2)
    check(abs(total - 600.5) < 1e-9, f'xls 求和 = 600.5（实际 {total}）')


def test_delete():
    """行/列删除：等于/包含匹配，.xlsx 与 .xls 两种格式。"""
    from openpyxl import Workbook as WX, load_workbook as LW
    clean_out()

    # --- 源 .xlsx ---
    src = OUT / 'd.xlsx'
    wb = WX()
    ws = wb.active
    ws['A1'] = '项目'
    ws['B1'] = '金额'
    ws['A2'] = '作废单'
    ws['B2'] = 10
    ws['A3'] = '正常单'
    ws['B3'] = 20
    ws['A4'] = '另一张'
    ws['B4'] = 30
    ws['A5'] = '作废2'
    ws['B5'] = 40
    ws['D1'] = '待删列'
    wb.save(src)

    def do_delete(mode, terms, case=False):
        plan, _ = excel_core.plan_delete([str(src)], mode, terms, OUT)
        excel_core.execute_delete(plan, mode, terms, case, log)
        return LW(plan[0].dst).worksheets[0]

    # 行：包含"作废" → 删第2、5行
    ws2 = do_delete('row', [('作废', 'contains')])
    colA = [ws2.cell(row=r, column=1).value for r in range(1, ws2.max_row + 1)]
    check(colA == ['项目', '正常单', '另一张'], f'行删除(包含作废) {colA}')

    # 行：等于"正常单" → 删第3行
    ws2 = do_delete('row', [('正常单', 'equals')])
    colA = [ws2.cell(row=r, column=1).value for r in range(1, ws2.max_row + 1)]
    check(colA == ['项目', '作废单', '另一张', '作废2'], f'行删除(等于) {colA}')

    # 多条件：同时删"作废"(包含) 和 "另一张"(等于) 两行
    ws2 = do_delete('row', [('作废', 'contains'), ('另一张', 'equals')])
    colA = [ws2.cell(row=r, column=1).value for r in range(1, ws2.max_row + 1)]
    check(colA == ['项目', '正常单'], f'多条件同时删除 {colA}')

    # 列：包含"待删" → 删 D 列
    ws2 = do_delete('column', [('待删', 'contains')])
    check(ws2.max_column == 2, f'列删除后 max_column={ws2.max_column}')
    check(ws2['A1'].value == '项目' and ws2['B1'].value == '金额', '列删除后保留 A、B 列')

    # 预览统计
    deleted, total = excel_core.count_delete(str(src), 'row', [('作废', 'contains')], False)
    check(deleted == 2 and total == 5, f'预览统计 删2行/共5行 ({deleted},{total})')

    # --- .xls ---
    import xlwt
    src_x = OUT / 'd.xls'
    wb = xlwt.Workbook()
    ws = wb.add_sheet('S')
    for r, v in enumerate(['项目', '作废单', '正常单', '作废2'], start=0):
        ws.write(r, 0, v)
    wb.save(str(src_x))
    plan, _ = excel_core.plan_delete([str(src_x)], 'row', [('作废', 'contains')], OUT)
    excel_core.execute_delete(plan, 'row', [('作废', 'contains')], False, log)
    import xlrd
    s = xlrd.open_workbook(str(plan[0].dst)).sheet_by_index(0)
    col = [s.cell_value(r, 0) for r in range(s.nrows)]
    check(col == ['项目', '正常单'], f'xls 行删除(包含作废) {col}')

    # --- 合并单元格处理：删除含合并单元格的行，合并区跟随上移 ---
    tpl = OUT / 'm.xlsx'
    wb = WX()
    ws = wb.active
    ws.merge_cells('B2:B3')
    ws['B2'] = '合并'
    ws['A1'] = 'h'
    ws['A3'] = '被删'
    wb.save(tpl)
    plan, _ = excel_core.plan_delete([str(tpl)], 'row', [('被删', 'contains')], OUT)
    excel_core.execute_delete(plan, 'row', [('被删', 'contains')], False, log)
    ws2 = LW(plan[0].dst).worksheets[0]
    check(ws2['B2'].value == '合并', '删除后合并内容保留')
    ranges = [str(r) for r in ws2.merged_cells.ranges]
    check('B2:B2' in ranges or 'B2' in [str(r) for r in ws2.merged_cells.ranges]
          or any('B2' in str(r) for r in ws2.merged_cells.ranges), f'合并区域调整 {ranges}')
    check(ws2['A2'].value is None or ws2['A2'].value != '被删', '被删行已移除')


def test_split():
    """端口拆分：分组、按流量总和大的一列降序、标红指定行、汇总表。"""
    from openpyxl import Workbook as WX, load_workbook as LW
    from openpyxl.styles import Border, Font, Side
    clean_out()

    src = OUT / '拆分源.xlsx'
    wb = WX()
    ws = wb.active
    ws.title = '数据'
    headers = ['集客系统端口号', '流入速度', '流出速度', '备注']
    for c, v in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=v)
    ws['A1'].font = Font(bold=True)

    rows = [
        ('P001', 10, 40, 'a'), ('P001', 20, 100, 'b'), ('P001', 30, 20, 'c'),
        ('P001', 40, 80, 'd'), ('P001', 50, 0, 'e'), ('P001', 60, 60, 'f'),
        ('P002', 5, 700, 'g'), ('P002', 15, 1000, 'h'), ('P002', 25, 800, 'i'),
        ('P002', 35, 900, 'j'),
        ('P003', 100, 1, 'k'), ('P003', 'N/A', 3, 'm'), ('P003', 200, 2, 'l'),
    ]
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 12
    ws.freeze_panes = 'A2'
    wb.save(src)

    # --- 分析（标红行用 5，避免生成 448 行）---
    a = excel_core.analyze_split(str(src), highlight_row=5)
    check(a.total_rows == 13, f'总数据行 13（实际 {a.total_rows}）')
    check([g.value for g in a.groups] == ['P001', 'P002', 'P003'],
          f'端口分组 {[g.value for g in a.groups]}')
    g1, g2, g3 = a.groups
    check(g1.sort_name == '流出速度' and g1.inflow_total == 210 and g1.outflow_total == 300,
          f'P001 按流出速度排序（{g1.inflow_total}/{g1.outflow_total}）')
    check(g2.sort_name == '流出速度' and g2.outflow_total == 3400, 'P002 按流出速度排序')
    check(g3.sort_name == '流入速度' and g3.inflow_total == 300, 'P003 按流入速度排序')
    check(g3.invalid_sort_values == 1, 'P003 非数值 1 个')
    check((g1.red_source_row, g2.red_source_row, g3.red_source_row) == (2, 8, None),
          f'标红来源行 {(g1.red_source_row, g2.red_source_row, g3.red_source_row)}')
    check(a.summary_rows == 2, f'汇总 2 行（实际 {a.summary_rows}）')

    prev = '\n'.join(excel_core.preview_split(str(src), highlight_row=5))
    check('P001' in prev and '流出速度' in prev and '标红并汇总' in prev, '预览文本包含关键信息')

    # 表头缺失时报错
    try:
        excel_core.analyze_split(str(src), port_header='不存在的列', highlight_row=5)
        check(False, '缺失列名应报错')
    except ValueError as e:
        check('未找到端口号列' in str(e), f'缺失列名报错：{e}')

    # --- 执行 ---
    plan, errors = excel_core.plan_split([str(src)], OUT)
    check(not errors and len(plan) == 1, '生成 1 个拆分计划')
    check(plan[0].dst.name == '拆分源_按端口拆分.xlsx', f'输出文件名 {plan[0].dst.name}')
    summary = excel_core.execute_split(
        plan, '集客系统端口号', '流入速度', '流出速度', 5, log)
    check(summary['ok'] == 1, f'执行成功 {summary}')

    out = LW(plan[0].dst)
    check(out.sheetnames == ['P001', 'P002', 'P003', '5行汇总'],
          f'工作表 {out.sheetnames}')
    for name in out.sheetnames:
        check([c.value for c in out[name][1]] == headers, f'{name} 表头与原表一致')

    remarks = lambda name: [out[name].cell(row=r, column=4).value
                            for r in range(2, out[name].max_row + 1)]
    check(remarks('P001') == ['b', 'd', 'f', 'a', 'c', 'e'],
          f'P001 按流出速度降序 {remarks("P001")}')
    check(remarks('P002') == ['h', 'j', 'i', 'g'],
          f'P002 按流出速度降序 {remarks("P002")}')
    check(remarks('P003') == ['l', 'k', 'm'],
          f'P003 按流入速度降序、非数值最后 {remarks("P003")}')

    # 第 5 行整行标红（含空列），其余数据行不红
    for name in ('P001', 'P002'):
        colors = [out[name].cell(row=5, column=c).font.color
                  for c in range(1, 5)]
        check(all(c is not None and c.rgb == 'FFFF0000' for c in colors),
              f'{name} 第 5 行整行标红')
    check(out['P003'].cell(row=2, column=1).font.color is None
          or out['P003'].cell(row=2, column=1).font.color.rgb != 'FFFF0000',
          'P003 数据不足，不标红')

    # 汇总表内容与标红行一致，且同样标红
    check([out['5行汇总'].cell(row=r, column=4).value for r in (2, 3)] == ['a', 'g'],
          '汇总表包含 P001/P002 的第 5 行')
    for r in (2, 3):
        colors = [out['5行汇总'].cell(row=r, column=c).font.color for c in range(1, 5)]
        check(all(c is not None and c.rgb == 'FFFF0000' for c in colors),
              f'汇总表第 {r} 行标红')

    # 样式/列宽/冻结窗格保留
    check(out['P001']['A1'].font.bold is True, '表头加粗样式保留')
    check(out['P001'].cell(row=2, column=1).border.left.style == 'thin', '数据边框保留')
    check(out['P001'].column_dimensions['A'].width == 18, '列宽保留')
    check(out['P001'].freeze_panes == 'A2', '冻结窗格保留')
    out.close()


def test_macos_filename_limits():
    """macOS 文件名按 UTF-8 字节截断（APFS 上限 255 字节）。"""
    old = sys.platform
    try:
        sys.platform = 'darwin'
        out = excel_core.truncate_stem('很' * 160)
        check(len(out.encode('utf-8')) <= 240,
              f'macOS 文件名 ≤240 字节（实际 {len(out.encode("utf-8"))}）')
        check(out == '很' * 80, f'macOS 中文长名截断到 {len(out)} 字')
        check(excel_core.truncate_stem('a' * 160) == 'a' * 150,
              'macOS ASCII 仍按 150 字符截断')
    finally:
        sys.platform = old


def test_config_paths():
    """配置目录：macOS 使用用户数据目录，Windows 保持软件目录优先。"""
    import config_store
    old = sys.platform
    try:
        sys.platform = 'darwin'
        mac = [str(p) for p in config_store.config_paths()]
        check(any('Application Support' in p and 'Lsy的Excel-Tool' in p for p in mac),
              f'macOS 配置目录 {mac}')
        sys.platform = 'win32'
        win = [str(p) for p in config_store.config_paths()]
        check(win and 'Lsy的Excel-Tool_profiles.json' in win[0],
              f'Windows 配置目录 {win}')
    finally:
        sys.platform = old


if __name__ == '__main__':
    test_rename()
    test_rename_advanced()
    test_fill()
    test_fill_column()
    test_fill_column_row_boundary()
    test_fill_column_merged()
    test_xls()
    test_fill_merged_cell()
    test_fill_sum()
    test_delete()
    test_split()
    test_macos_filename_limits()
    test_config_paths()
    print('\n全部测试通过 ✅')
