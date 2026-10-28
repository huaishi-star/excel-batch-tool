"""Excel 批量处理核心逻辑（GUI 无关，可独立测试）。

提供以下能力：
1. 批量重命名：按"单元格组合"规则生成新文件名，复制到输出文件夹（原文件不动）
2. 模板填充：把源文件的单元格值按映射关系填入模板对应单元格，保留模板格式
3. 批量行/列删除：按条件删除整行/整列
4. 端口拆分：按某一列的不同值拆成多个工作表，按流量列排序并标红指定行、汇总标红行

注意：
- 单元格定位一律使用文件/模板的第一个工作表（wb.worksheets[0]）
- 读取使用 data_only=True, read_only=True；模板填充写入使用完整模式以保留格式
- 端口拆分读取使用 read_only + write_only，避免大表（每个端口约 8928 行）占用过多内存
"""

from __future__ import annotations

import re
import shutil
import sys
import xml.etree.ElementTree as ET
import zipfile
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from threading import Event
from typing import Callable, Iterable, Optional

import xlrd
import xlwt
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

# ---------------------------------------------------------------------------
# 常量与正则
# ---------------------------------------------------------------------------

# Windows 文件名非法字符
ILLEGAL_CHARS_RE = re.compile(r'[<>:"/\\|?*]')
# 控制字符（换行、制表等）
CONTROL_CHARS_RE = re.compile(r'[\x00-\x1f]')
# Windows 保留设备名
RESERVED_NAMES = (
    {'CON', 'PRN', 'AUX', 'NUL'}
    | {f'COM{i}' for i in range(1, 10)}
    | {f'LPT{i}' for i in range(1, 10)}
)
# 单元格坐标，如 A1、AB12
CELL_RE = re.compile(r'^[A-Za-z]{1,3}[0-9]{1,7}$')
# 列字母，如 A、B、AA
COL_RE = re.compile(r'^[A-Za-z]{1,3}$')
# 坐标匹配：列字母 + 行号
_COORD_MATCH = re.compile(r'^([A-Za-z]+)([0-9]+)$')

# 整列读取的行数上限（安全阀）
MAX_COL_ROWS = 5000


def _col_letter_to_idx(col: str) -> int:
    """列字母 A、B、AA → 0 基列号。"""
    idx = 0
    for ch in col:
        idx = idx * 26 + (ord(ch.upper()) - 64)
    return idx - 1


def _coord_to_rc(coord: str) -> tuple:
    """'A1' → (0 基行号, 0 基列号)。"""
    m = _COORD_MATCH.match(coord)
    return int(m.group(2)) - 1, _col_letter_to_idx(m.group(1))


def _openpyxl_merged_anchor(ws, coord: str) -> str:
    """若 coord 位于某个合并区域内，返回其左上角(锚点)坐标；否则原样返回。

    仅用于"单元格→单元格"映射：向合并区写入时把值放到锚点，保证显示。
    """
    col_letter, row = coordinate_from_string(coord)
    col = column_index_from_string(col_letter)
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
            return f'{get_column_letter(mr.min_col)}{mr.min_row}'
    return coord


def _openpyxl_merged_range(ws, coord: str):
    """若 coord 位于某个合并区域内，返回该区域对象；否则返回 None。"""
    col_letter, row = coordinate_from_string(coord)
    col = column_index_from_string(col_letter)
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
            return mr
    return None


def _shift_openpyxl_merges(ws, from_row: int, amount: int):
    """插入行后，把 from_row 及以下的合并区域整体下移 amount 行（跨插入点的区域向下扩展）。"""
    if amount <= 0:
        return
    for mr in ws.merged_cells.ranges:
        if mr.min_row >= from_row:
            mr.shift(0, amount)
        elif mr.min_row < from_row <= mr.max_row:
            mr.max_row += amount


def _xls_merged_anchor(merged: list, r: int, c: int) -> tuple:
    """若 (r,c) 位于某个合并区域内（0 基），返回其左上角(锚点)坐标；否则原样返回。"""
    for (rlo, rhi, clo, chi) in merged:
        if rlo <= r < rhi and clo <= c < chi:
            return rlo, clo
    return r, c


def _xlrd_cell_value(book, sheet, r: int, c: int):
    """读取 xlrd 单元格为 Python 值：日期→datetime、布尔→bool、空→None。

    越界（r/c 超出工作表使用范围）按空单元格处理，
    避免 xlrd 对越界访问抛 IndexError（openpyxl 则默认返回空）。
    """
    if r >= sheet.nrows or c >= sheet.ncols:
        return None
    cell = sheet.cell(r, c)
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    return cell.value

# 文件名主干最长字符数（Windows 单段建议 ≤255，留余量）
MAX_STEM = 150
# 完整路径最长字符数（避开 Windows MAX_PATH=260，留余量）
MAX_PATH = 250
# macOS 完整路径上限（PATH_MAX=1024 字节，留余量给重名后缀）
MAX_PATH_BYTES_MAC = 980

LogFn = Callable[[str], None]
ProgressFn = Callable[[int, int], None]


# ---------------------------------------------------------------------------
# 数据类
# ---------------------------------------------------------------------------

@dataclass
class PlanItem:
    """重命名计划中的一项。"""
    src: Path
    dst: Path
    name: str  # dst 的文件名，供预览


@dataclass
class FillItem:
    """模板填充计划中的一项。"""
    src: Path
    dst: Path


@dataclass
class DeleteItem:
    """行/列删除计划中的一项。"""
    src: Path
    dst: Path


# ---------------------------------------------------------------------------
# 单元格值 → 文本
# ---------------------------------------------------------------------------

def value_to_text(value, default: str = "") -> str:
    """把单元格值转换成适合用于文件名的显示文本。

    - None             → default（通常为空串，拼接时会被跳过）
    - datetime / date  → %Y%m%d（如 20240115，避免出现 / 或 - 等非法/歧义字符）
    - bool             → TRUE / FALSE（与 Excel 显示一致）
    - 整型浮点          → 去掉 .0（98.0 → '98'；98.5 → '98.5'）
    - 其余              → str(value)
    """
    if value is None:
        return default
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y%m%d')
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_cell_text(path, coord: str) -> str:
    """读取文件第一个工作表中 coord 单元格的显示文本（供重命名用），支持 .xls/.xlsx/.xlsm。"""
    if Path(path).suffix.lower() == '.xls':
        book = xlrd.open_workbook(str(path))
        try:
            r, c = _coord_to_rc(coord)
            return value_to_text(_xlrd_cell_value(book, book.sheet_by_index(0), r, c))
        finally:
            book.release_resources()
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return value_to_text(wb.worksheets[0][coord].value)
    finally:
        wb.close()


def read_cell_raw(path, coord):
    """读取文件第一个工作表中 coord 单元格的原始 Python 值（供模板填充用）。

    None 会转成空串，以便写入时清空模板对应单元格。支持 .xls/.xlsx/.xlsm。
    """
    if Path(path).suffix.lower() == '.xls':
        book = xlrd.open_workbook(str(path))
        try:
            r, c = _coord_to_rc(coord)
            v = _xlrd_cell_value(book, book.sheet_by_index(0), r, c)
            return '' if v is None else v
        finally:
            book.release_resources()
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        v = wb.worksheets[0][coord].value
        return '' if v is None else v
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# 文件名安全处理
# ---------------------------------------------------------------------------

def sanitize_name_part(part: str) -> str:
    """净化文件名片段：
    - 去除控制字符
    - 非法字符替换为下划线
    - 去掉首尾空格与点号
    - 若是 Windows 保留设备名，加下划线前缀
    """
    s = CONTROL_CHARS_RE.sub('', part)
    s = ILLEGAL_CHARS_RE.sub('_', s)
    s = s.strip(' .')
    if not s:
        return s
    if s.split('.')[0].upper() in RESERVED_NAMES:
        s = '_' + s
    return s


def truncate_stem(stem: str, max_stem: int = MAX_STEM) -> str:
    """把主文件名截断到 max_stem 字符，并保证结尾不是点或空格。

    macOS 的 APFS 文件名上限是 255 个 UTF-8 字节，中文名按字符截断后仍可能超限，
    因此在 macOS 上再按字节截断一次（240 字节，留出扩展名余量）。
    """
    if len(stem) > max_stem:
        stem = stem[:max_stem]
    if sys.platform == 'darwin':
        raw = stem.encode('utf-8')
        if len(raw) > 240:
            stem = raw[:240].decode('utf-8', errors='ignore')
    return stem.rstrip(' .')


def unique_path(dst: Path, used: set[str]) -> Path:
    """目标路径若已存在（磁盘上或本次任务中已用过），自动追加 (1)、(2)…"""
    if dst.name.lower() not in used and not dst.exists():
        return dst
    stem, ext = dst.stem, dst.suffix
    n = 1
    while True:
        candidate = dst.with_name(f'{stem} ({n}){ext}')
        if candidate.name.lower() not in used and not candidate.exists():
            return candidate
        n += 1


def build_dest_name(stem: str, ext: str, out_dir, used: set[str],
                    max_stem: int = MAX_STEM, max_path: int = MAX_PATH) -> Path:
    """根据文件名主干和扩展名构造输出路径。

    依次处理：净化 → 截断 → 保证完整路径长度 → 重名去重。
    """
    stem = truncate_stem(sanitize_name_part(stem), max_stem)
    if not stem:
        stem = '未命名'
    dst = Path(out_dir) / f'{stem}{ext}'
    # 完整路径过长时，逐段缩短主干
    while len(str(dst)) > max_path and len(stem) > 5:
        stem = stem[:-5].rstrip(' .')
        dst = Path(out_dir) / f'{stem}{ext}'
    # macOS 的 PATH_MAX 按字节计算，中文路径可能字符数没超但字节数超了
    if sys.platform == 'darwin':
        while len(str(dst).encode('utf-8')) > MAX_PATH_BYTES_MAC and len(stem) > 5:
            stem = stem[:-5].rstrip(' .')
            dst = Path(out_dir) / f'{stem}{ext}'
    return unique_path(dst, used)


# ---------------------------------------------------------------------------
# 批量重命名
# ---------------------------------------------------------------------------

def build_name(parts: list[tuple], separator: str) -> str:
    """按规则元素拼接文件名主干。

    parts: [(kind, value)]，kind ∈ {'text', 'data'}。
    - 'text'：手打文字，原样保留；
    - 'data'：单元格内容或原文件名，仅当非空时加入；且只在相邻两个 data 之间插入分隔符。
    """
    out: list[str] = []
    prev_data = False
    for kind, value in parts:
        if kind == 'text':
            out.append(value)
            prev_data = False
        else:  # data
            if value:
                if prev_data and separator:
                    out.append(separator)
                out.append(value)
                prev_data = True
    return ''.join(out)


def plan_rename(files: Iterable, rule: list, separator: str, out_dir,
                ) -> tuple[list[PlanItem], list[tuple]]:
    """生成重命名计划（不执行）。

    rule 是规则元素列表，每个元素为 (kind, value)：
      ('cell', 'A1')     → 读取文件第一个工作表该单元格的显示文本
      ('text', '文字')   → 原样保留的手打文字
      ('filename', '')   → 原文件名（不含扩展名）
    元素按列表顺序拼接；分隔符只插在相邻两个"单元格/原文件名"之间，文字原样保留。
    所有内容都为空时回退到原文件名。返回 (计划列表, 错误列表)。
    """
    used: set[str] = set()
    plan: list[PlanItem] = []
    errors: list[tuple] = []
    for f in files:
        f = Path(f)
        try:
            parts: list[tuple] = []
            for kind, value in rule:
                if kind == 'text':
                    parts.append(('text', value))
                elif kind == 'filename':
                    parts.append(('data', f.stem))
                else:  # cell
                    parts.append(('data', read_cell_text(f, value)))
        except Exception as e:  # 文件损坏 / 被占用等
            errors.append((f, str(e)))
            continue
        stem = build_name(parts, separator)
        if not stem:
            stem = f.stem  # 全部为空时回退原文件名
        ext = f.suffix.lower()
        dst = build_dest_name(stem, ext, out_dir, used)
        used.add(dst.name.lower())
        plan.append(PlanItem(src=f, dst=dst, name=dst.name))
    return plan, errors


def execute_rename(plan: list[PlanItem], log: LogFn, progress: Optional[ProgressFn] = None,
                   cancel_event: Optional[Event] = None) -> dict:
    """执行重命名计划：把原文件复制到目标路径（shutil.copy2 保留 .xlsm 宏）。

    单文件失败不中断批次。返回 {'ok': n, 'skip': n, 'fail': n}。
    """
    if not plan:
        return {'ok': 0, 'skip': 0, 'fail': 0}
    plan[0].dst.parent.mkdir(parents=True, exist_ok=True)
    total = len(plan)
    ok = skip = fail = 0
    for i, item in enumerate(plan):
        if cancel_event is not None and cancel_event.is_set():
            log('[取消] 用户取消操作')
            break
        try:
            shutil.copy2(item.src, item.dst)
            ok += 1
            log(f'[OK] {item.src.name} → {item.dst.name}')
        except PermissionError:
            skip += 1
            log(f'[跳过] {item.src.name}：文件被占用，已跳过')
        except Exception as e:
            fail += 1
            log(f'[失败] {item.src.name}：{e}')
        if progress is not None:
            progress(i + 1, total)
    return {'ok': ok, 'skip': skip, 'fail': fail}


# ---------------------------------------------------------------------------
# 模板填充
# ---------------------------------------------------------------------------

def plan_fill(src_files: Iterable, template_path, mapping: list[tuple], out_dir,
              ) -> tuple[list[FillItem], list[tuple]]:
    """生成模板填充计划（不执行）。

    输出文件名 = 模板名 + '_' + 源文件名（保持模板扩展名），自动去重。
    返回 (计划列表, 错误列表)。
    """
    template_path = Path(template_path)
    tstem = template_path.stem
    text = template_path.suffix.lower()
    used: set[str] = set()
    plan: list[FillItem] = []
    errors: list[tuple] = []
    for f in src_files:
        f = Path(f)
        try:
            dst = build_dest_name(f'{tstem}_{f.stem}', text, out_dir, used)
        except Exception as e:
            errors.append((f, str(e)))
            continue
        used.add(dst.name.lower())
        plan.append(FillItem(src=f, dst=dst))
    return plan, errors


def _copy_row_style(ws, src_row: int, dst_row: int):
    """把 src_row 的样式（字体/边框/填充/数字格式/对齐）复制到 dst_row。

    用于让插入的空白行延续下方数据行的表格边框。
    """
    for col_idx in range(1, ws.max_column + 1):
        s = ws.cell(row=src_row, column=col_idx)
        d = ws.cell(row=dst_row, column=col_idx)
        d.font = copy(s.font)
        d.border = copy(s.border)
        d.fill = copy(s.fill)
        d.number_format = s.number_format
        d.alignment = copy(s.alignment)


def _fill_column(ws, values: list, col: str, start_row: int) -> int:
    """把 values 写入模板 ws 第 col 列，从 start_row 行开始。

    先在目标列内找可用空白（自 start_row 向下，遇到目标列已有内容或合并单元格为止）；
    若源数据超出空白，在边界处一次性批量插入差额空白行（把边界及下方内容、合并区域
    整体下移），绝不覆盖、也绝不因其他列有数据而影响对齐。
    返回插入的总行数。
    """
    # 1) 计算目标列可用空白与边界行（合并单元格视为已占用）
    boundary = None
    available = 0
    row = start_row
    while row <= ws.max_row and available < len(values):
        coord = f'{col}{row}'
        cell = ws[coord]
        if (cell.value is not None and str(cell.value).strip() != '') \
                or _openpyxl_merged_range(ws, coord) is not None:
            boundary = row
            break
        available += 1
        row += 1

    # 2) 源数据超出空白时，在边界处批量插入差额空白行
    shortfall = max(0, len(values) - available) if boundary is not None else 0
    if shortfall:
        ws.insert_rows(boundary, shortfall)
        _shift_openpyxl_merges(ws, boundary, shortfall)
        # 让插入的空行继承下方(边界)行的样式，表格边框连续
        for i in range(shortfall):
            _copy_row_style(ws, boundary + shortfall, boundary + i)

    # 3) 直接写入目标单元格（合并区已在扫描中作为边界排除，避免覆盖合并内容）
    for i, v in enumerate(values):
        r = start_row + i
        ws[f'{col}{r}'] = v
    return shortfall


# ---------------------------------------------------------------------------
# .xls（旧版 Excel）支持：读取 / 模板填充（含插入空白行）
# ---------------------------------------------------------------------------

class _XlsSource:
    """.xls 源文件读取器（xlrd）。"""

    def __init__(self, path):
        self.book = xlrd.open_workbook(str(path))
        self.sheet = self.book.sheet_by_index(0)

    def cell(self, coord):
        r, c = _coord_to_rc(coord)
        v = _xlrd_cell_value(self.book, self.sheet, r, c)
        return '' if v is None else v

    def column(self, col, start_row):
        c = _col_letter_to_idx(col)
        r = start_row - 1
        values = []
        while r < start_row - 1 + MAX_COL_ROWS:
            v = _xlrd_cell_value(self.book, self.sheet, r, c)
            if v is None or (isinstance(v, str) and v == ''):
                break
            values.append(v)
            r += 1
        return values

    def close(self):
        self.book.release_resources()


class _OpenpyxlSource:
    """.xlsx/.xlsm 源文件读取器（openpyxl）。"""

    def __init__(self, path):
        self.wb = load_workbook(path, data_only=True, read_only=True)
        self.ws = self.wb.worksheets[0]

    def cell(self, coord):
        v = self.ws[coord].value
        return '' if v is None else v

    def column(self, col, start_row):
        values = []
        r = start_row
        while r < start_row + MAX_COL_ROWS:
            v = self.ws[f'{col}{r}'].value
            if v is None or (isinstance(v, str) and v == ''):
                break
            values.append(v)
            r += 1
        return values

    def close(self):
        self.wb.close()


def _open_source(path):
    if Path(path).suffix.lower() == '.xls':
        return _XlsSource(path)
    return _OpenpyxlSource(path)


# ---------------------------------------------------------------------------
# 列求和：把源文件某列（含数值型字符串、前后空格）准确转成数字并求和
# ---------------------------------------------------------------------------

def _to_number(value):
    """把单元格值转换为数字：数值直接返回；字符串去掉首尾空格（含千分位逗号）后转换。

    无法转换（空、布尔、非数字文本）返回 None。
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            try:
                return float(s.replace(',', ''))
            except ValueError:
                return None
    return None


def _sum_values(values) -> tuple:
    """对值列表中的数值（含数值型字符串）求和。返回 (总和, 有效数值个数)。"""
    total = 0.0
    count = 0
    for v in values:
        n = _to_number(v)
        if n is not None:
            total += n
            count += 1
    return total, count


def sum_column(path, col: str, start_row: int) -> tuple:
    """读取源文件某列自 start_row 起（到第一个空单元格止）的所有数值并求和。

    返回 (总和, 有效数值个数)，供预览与执行共用。
    """
    src = _open_source(path)
    try:
        values = src.column(col, start_row)
    finally:
        src.close()
    return _sum_values(values)


# Excel 内置数字格式串（xlrd 对部分内置格式返回 format_str=None，需自行补齐）
_BUILTIN_FORMATS = {
    1: '0', 2: '0.00', 3: '#,##0', 4: '#,##0.00',
    5: '$#,##0_);($#,##0)', 6: '$#,##0_);[Red]($#,##0)',
    7: '$#,##0.00_);($#,##0.00)', 8: '$#,##0.00_);[Red]($#,##0.00)',
    9: '0%', 10: '0.00%', 11: '0.00E+00', 12: '# ?/?', 13: '# ??/??',
    14: 'm/d/yy', 15: 'd-mmm-yy', 16: 'd-mmm', 17: 'mmm-yy',
    18: 'h:mm AM/PM', 19: 'h:mm:ss AM/PM', 20: 'h:mm', 21: 'h:mm:ss',
    22: 'm/d/yy h:mm', 49: '@',
}


def _resolve_num_format(book, format_key):
    """取数字格式串；格式串为 None 时回退到常用内置格式，避免 xlwt 对 None 报错。"""
    fmt = book.format_map.get(format_key)
    if fmt is not None and fmt.format_str:
        return fmt.format_str
    return _BUILTIN_FORMATS.get(format_key, '')


def _make_xlwt_style(book, xf_index):
    """把 xlrd 的 XF 格式转换为 xlwt 样式，用于写入时保留原格式。"""
    style = xlwt.XFStyle()
    if xf_index is None:
        return style
    xf = book.xf_list[xf_index]

    font_rec = book.font_list[xf.font_index]
    f = xlwt.Font()
    f.name = font_rec.name
    f.bold = bool(font_rec.bold)
    f.italic = bool(font_rec.italic)
    f.height = font_rec.height
    f.colour_index = font_rec.colour_index
    style.font = f

    b = xlwt.Borders()
    b.left = xf.border.left_line_style
    b.right = xf.border.right_line_style
    b.top = xf.border.top_line_style
    b.bottom = xf.border.bottom_line_style
    style.borders = b

    pat = xlwt.Pattern()
    pat.pattern = xf.background.fill_pattern
    pat.pattern_fore_colour = xf.background.pattern_colour_index
    pat.pattern_back_colour = xf.background.background_colour_index
    style.pattern = pat

    al = xlwt.Alignment()
    al.horz = xf.alignment.hor_align
    al.vert = xf.alignment.vert_align
    al.wrap = 1 if xf.alignment.text_wrapped else 0
    style.alignment = al

    style.num_format_str = _resolve_num_format(book, xf.format_key) or 'General'
    return style


def _shift_rows_down(cells: dict, styles: dict, merged: list, row: int):
    """把 cells/styles 中所有行号 >= row 的单元格下移一行（0 基），并同步合并单元格。"""
    keys = sorted([k for k in cells if k[0] >= row], key=lambda k: k[0], reverse=True)
    for k in keys:
        cells[(k[0] + 1, k[1])] = cells.pop(k)
        styles[(k[0] + 1, k[1])] = styles.pop(k)
    new_merged = []
    for m in merged:
        rlo, rhi, clo, chi = m
        if rlo >= row:
            new_merged.append((rlo + 1, rhi + 1, clo, chi))
        elif rhi > row:  # 合并范围跨过插入点 → 向下扩展
            new_merged.append((rlo, rhi + 1, clo, chi))
        else:
            new_merged.append((rlo, rhi, clo, chi))
    merged[:] = new_merged


def _fill_xls_template(template_path, cell_values: dict, column_jobs: list, output_path):
    """把单元格/整列数据填入 .xls 模板并另存为 .xls。

    .xls 为旧二进制格式，无插入行能力，因此采用"整表重建"：
    用 xlrd 读取原表全部内容与格式 → 计算插入行与写入 → 用 xlwt 重建保存。
    目标单元格已有内容时插入空白行下移，绝不覆盖。
    """
    book = xlrd.open_workbook(str(template_path), formatting_info=True)
    src = book.sheet_by_index(0)

    # 载入原表：cells[(r,c)]=值，styles[(r,c)]=xf_index（保留有值或有格式的单元格）
    cells: dict = {}
    styles: dict = {}
    for r in range(src.nrows):
        for c in range(src.ncols):
            v = src.cell_value(r, c)
            xf = src.cell_xf_index(r, c)
            if v != '' or xf != 0:
                cells[(r, c)] = v
                styles[(r, c)] = xf
    merged = [list(m) for m in src.merged_cells]

    # 单元格映射（目标在合并区内时重定向到锚点，保证值可见）
    for tpl_coord, value in cell_values.items():
        r, c = _coord_to_rc(tpl_coord)
        r, c = _xls_merged_anchor(merged, r, c)
        cells[(r, c)] = value
        styles.setdefault((r, c), None)

    # 整列映射（源数据超出空白时在边界处批量插入，不因其他列有数据而影响对齐）
    for tpl_col, tpl_start, values in column_jobs:
        c = _col_letter_to_idx(tpl_col)
        row = tpl_start - 1
        # 计算目标列可用空白与边界行（合并单元格视为已占用）
        boundary = None
        available = 0
        probe = row
        while available < len(values):
            occupied = False
            v = cells.get((probe, c))
            if v is not None and str(v).strip() != '':
                occupied = True
            else:
                for (rlo, rhi, clo, chi) in merged:
                    if rlo <= probe < rhi and clo <= c < chi:
                        occupied = True
                        break
            if occupied:
                boundary = probe
                break
            available += 1
            probe += 1
        # 在边界处批量下移差额行（_shift_rows_down 会同步移动合并区域）
        shortfall = max(0, len(values) - available) if boundary is not None else 0
        for _ in range(shortfall):
            _shift_rows_down(cells, styles, merged, boundary)
        # 直接写入目标单元格（不重定向到合并锚点，避免覆盖合并内容）
        for v in values:
            cells[(row, c)] = v
            if (row, c) not in styles:
                styles[(row, c)] = styles.get((row - 1, c))
            row += 1

    # 重建 .xls
    wb = xlwt.Workbook()
    ws = wb.add_sheet(src.name or 'Sheet1')
    # 合并区域内的单元格（含左上角）统一交给 write_merge 处理，
    # 逐格写入时跳过，以免 xlwt 抛 "Attempt to overwrite cell"
    covered = set()
    for (rlo, rhi, clo, chi) in merged:
        for r in range(rlo, rhi):
            for c in range(clo, chi):
                covered.add((r, c))
    for (r, c), v in cells.items():
        if (r, c) in covered:
            continue
        ws.write(r, c, v, _make_xlwt_style(book, styles.get((r, c))))
    for (rlo, rhi, clo, chi) in merged:
        ws.write_merge(rlo, rhi - 1, clo, chi - 1, cells.get((rlo, clo), ''),
                       _make_xlwt_style(book, styles.get((rlo, clo))))
    for col_idx, colinfo in src.colinfo_map.items():
        ws.col(col_idx).width = colinfo.width
    for r, rowinfo in src.rowinfo_map.items():
        if rowinfo.height:
            ws.row(r).height = rowinfo.height
    wb.save(str(output_path))
    book.release_resources()


def execute_fill(template_path, plan: list[FillItem], mapping: list[tuple],
                 log: LogFn, progress: Optional[ProgressFn] = None,
                 cancel_event: Optional[Event] = None) -> dict:
    """执行模板填充计划。

    映射条目支持三种（均指第一个工作表）：
      ('cell', 'A1', 'B2')      → 源单元格内容写入模板单元格（None 写空串以清空）
      ('column', 'C', 2, 'B', 6)→ 源文件 C 列自第 2 行向下（到空单元格为止）的数据，
                                 写入模板 B 列自第 6 行开始；目标已有内容时插入空白行下移，不覆盖
      ('sum', 'B', 2, 'D5')     → 源文件 B 列自第 2 行向下（到空单元格为止）的所有数值
                                 求和，结果写入模板 D5（支持合并单元格，自动定位锚点）

    每个源文件：重新加载模板副本（避免输出互相污染），写入后保存到目标路径。
    - .xlsx/.xlsm 模板：openpyxl 往返保留格式；整列映射用插入空白行扩展。
    - .xls 模板：整表重建（xlrd 读 + xlwt 写），同样支持插入空白行。
    返回 {'ok': n, 'skip': n, 'fail': n}。
    """
    if not plan:
        return {'ok': 0, 'skip': 0, 'fail': 0}
    template_path = Path(template_path)
    is_xls_tpl = template_path.suffix.lower() == '.xls'
    keep_vba = template_path.suffix.lower() == '.xlsm'
    plan[0].dst.parent.mkdir(parents=True, exist_ok=True)
    total = len(plan)
    ok = skip = fail = 0
    for i, item in enumerate(plan):
        if cancel_event is not None and cancel_event.is_set():
            log('[取消] 用户取消操作')
            break
        wb = None
        src = None
        try:
            # 1) 读取源文件数据（每个源文件只打开一次，.xls/.xlsx/.xlsm 均可）
            cell_values: dict = {}   # 模板坐标 → 值
            column_jobs: list = []   # (模板列, 起始行, 数据列表)
            src = _open_source(item.src)
            for m in mapping:
                if m[0] == 'cell':
                    cell_values[m[2]] = src.cell(m[1])
                elif m[0] == 'column':
                    column_jobs.append((m[3], m[4], src.column(m[1], m[2])))
                elif m[0] == 'sum':
                    _, src_col, src_start, tpl_coord = m
                    total, _ = _sum_values(src.column(src_col, src_start))
                    cell_values[tpl_coord] = total
            src.close()
            src = None

            # 2) 写入模板副本（先写单元格映射，再写整列映射，可能插入行）
            if is_xls_tpl:
                _fill_xls_template(template_path, cell_values, column_jobs, item.dst)
            else:
                wb = load_workbook(template_path, keep_vba=keep_vba)
                ws = wb.worksheets[0]
                for tpl_coord, v in cell_values.items():
                    # 目标在合并区内时重定向到锚点，保证值在合并单元格可见
                    ws[_openpyxl_merged_anchor(ws, tpl_coord)] = v
                for tpl_col, tpl_start, values in column_jobs:
                    _fill_column(ws, values, tpl_col, tpl_start)
                wb.save(item.dst)
                wb.close()
                wb = None
            ok += 1
            log(f'[OK] {item.src.name} → {item.dst.name}')
        except PermissionError:
            skip += 1
            log(f'[跳过] {item.src.name}：文件被占用，已跳过')
        except Exception as e:
            fail += 1
            log(f'[失败] {item.src.name}：{e}')
        finally:
            for _w in (wb, src):
                if _w is not None:
                    try:
                        _w.close()
                    except Exception:
                        pass
        if progress is not None:
            progress(i + 1, total)
    return {'ok': ok, 'skip': skip, 'fail': fail}


# ---------------------------------------------------------------------------
# 批量行/列删除
# ---------------------------------------------------------------------------

def _cell_matches(value, target, match_type: str, case_sensitive: bool) -> bool:
    """判断单元格值是否满足单条删除条件：'contains' 包含 / 'equals' 等于。"""
    if value is None:
        return False
    s = value_to_text(value)
    t = str(target)
    if match_type == 'contains':
        return (t in s) if case_sensitive else (t.lower() in s.lower())
    return (s.strip() == t.strip()) if case_sensitive else (s.strip().lower() == t.strip().lower())


def _cell_matches_any(value, terms: list, case_sensitive: bool) -> bool:
    """判断单元格值是否满足任意一条条件。terms = [(值, 匹配方式), ...]。"""
    if value is None:
        return False
    for target, match_type in terms:
        if _cell_matches(value, target, match_type, case_sensitive):
            return True
    return False


def _rows_to_delete(ws, terms: list, case_sensitive: bool) -> set:
    """扫描工作表中"任意单元格满足任意条件"的行。"""
    deleted = set()
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if _cell_matches_any(ws.cell(row=row, column=col).value, terms, case_sensitive):
                deleted.add(row)
                break
    return deleted


def _cols_to_delete(ws, terms: list, case_sensitive: bool) -> set:
    """扫描工作表中"任意单元格满足任意条件"的列。"""
    deleted = set()
    for col in range(1, ws.max_column + 1):
        for row in range(1, ws.max_row + 1):
            if _cell_matches_any(ws.cell(row=row, column=col).value, terms, case_sensitive):
                deleted.add(col)
                break
    return deleted


def _fix_merges_after_row_delete(ws, deleted_rows: set):
    """删除行后同步修正合并单元格：整块被删的移除，其余按上方删除数上移。"""
    for mr in list(ws.merged_cells.ranges):
        new_min = mr.min_row - sum(1 for r in deleted_rows if r < mr.min_row)
        new_max = mr.max_row - sum(1 for r in deleted_rows if r <= mr.max_row)
        if new_min > new_max:
            ws.unmerge_cells(str(mr))
        else:
            mr.min_row = new_min
            mr.max_row = new_max


def _fix_merges_after_col_delete(ws, deleted_cols: set):
    """删除列后同步修正合并单元格。"""
    for mr in list(ws.merged_cells.ranges):
        new_min = mr.min_col - sum(1 for c in deleted_cols if c < mr.min_col)
        new_max = mr.max_col - sum(1 for c in deleted_cols if c <= mr.max_col)
        if new_min > new_max:
            ws.unmerge_cells(str(mr))
        else:
            mr.min_col = new_min
            mr.max_col = new_max


def _delete_openpyxl_rows(ws, deleted_rows: set):
    """按连续块自下而上删除行。"""
    _fix_merges_after_row_delete(ws, deleted_rows)
    rows = sorted(deleted_rows)
    i = len(rows) - 1
    while i >= 0:
        end = i
        while i > 0 and rows[i] - rows[i - 1] == 1:
            i -= 1
        start = i
        ws.delete_rows(rows[start], end - start + 1)
        i -= 1


def _delete_openpyxl_cols(ws, deleted_cols: set):
    """按连续块自右向左删除列。"""
    _fix_merges_after_col_delete(ws, deleted_cols)
    cols = sorted(deleted_cols)
    i = len(cols) - 1
    while i >= 0:
        end = i
        while i > 0 and cols[i] - cols[i - 1] == 1:
            i -= 1
        start = i
        ws.delete_cols(cols[start], end - start + 1)
        i -= 1


def _delete_openpyxl(src, dst, mode: str, terms: list, case_sensitive: bool):
    """用 openpyxl 删除 .xlsx/.xlsm 中的行/列并保存。"""
    wb = load_workbook(str(src))
    try:
        ws = wb.worksheets[0]
        if mode == 'row':
            _delete_openpyxl_rows(ws, _rows_to_delete(ws, terms, case_sensitive))
        else:
            _delete_openpyxl_cols(ws, _cols_to_delete(ws, terms, case_sensitive))
        wb.save(str(dst))
    finally:
        wb.close()


def _xls_deleted_rows_cols(sheet, mode: str, terms: list, case_sensitive: bool) -> tuple:
    """统计 .xls 中要删除的行/列。"""
    if mode == 'row':
        deleted = set()
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                if _cell_matches_any(sheet.cell_value(r, c), terms, case_sensitive):
                    deleted.add(r)
                    break
        return deleted, set()
    deleted = set()
    for c in range(sheet.ncols):
        for r in range(sheet.nrows):
            if _cell_matches_any(sheet.cell_value(r, c), terms, case_sensitive):
                deleted.add(c)
                break
    return set(), deleted


def _rebuild_xls_deleted(book, src_sheet, deleted_rows: set, deleted_cols: set, output_path):
    """跳过被删除的行/列，用 xlwt 重建 .xls 文件。"""
    row_map = {}
    new_r = 0
    for r in range(src_sheet.nrows):
        if r not in deleted_rows:
            row_map[r] = new_r
            new_r += 1
    col_map = {}
    new_c = 0
    for c in range(src_sheet.ncols):
        if c not in deleted_cols:
            col_map[c] = new_c
            new_c += 1
    wb = xlwt.Workbook()
    ws = wb.add_sheet(src_sheet.name or 'Sheet1')
    for r in range(src_sheet.nrows):
        if r in deleted_rows:
            continue
        for c in range(src_sheet.ncols):
            if c in deleted_cols:
                continue
            ws.write(row_map[r], col_map[c], src_sheet.cell_value(r, c),
                     _make_xlwt_style(book, src_sheet.cell_xf_index(r, c)))
    # 合并单元格：左上角被删除则丢弃，否则按保留行列收缩
    for (rlo, rhi, clo, chi) in src_sheet.merged_cells:
        if rlo not in row_map or clo not in col_map:
            continue
        last_r = rhi - 1
        while last_r >= rlo and last_r not in row_map:
            last_r -= 1
        last_c = chi - 1
        while last_c >= clo and last_c not in col_map:
            last_c -= 1
        if last_r < rlo or last_c < clo:
            continue
        ws.write_merge(row_map[rlo], row_map[last_r], col_map[clo], col_map[last_c],
                       src_sheet.cell_value(rlo, clo),
                       _make_xlwt_style(book, src_sheet.cell_xf_index(rlo, clo)))
    for col_idx, colinfo in src_sheet.colinfo_map.items():
        if col_idx in col_map:
            ws.col(col_map[col_idx]).width = colinfo.width
    for r, rowinfo in src_sheet.rowinfo_map.items():
        if r in row_map and rowinfo.height:
            ws.row(row_map[r]).height = rowinfo.height
    wb.save(str(output_path))


def _delete_xls(src, dst, mode: str, terms: list, case_sensitive: bool):
    """删除 .xls 中的行/列并重建保存。"""
    book = xlrd.open_workbook(str(src), formatting_info=True)
    try:
        sheet = book.sheet_by_index(0)
        deleted_rows, deleted_cols = _xls_deleted_rows_cols(sheet, mode, terms, case_sensitive)
        _rebuild_xls_deleted(book, sheet, deleted_rows, deleted_cols, dst)
    finally:
        book.release_resources()


def plan_delete(files: Iterable, mode: str, terms: list, out_dir,
                ) -> tuple[list[DeleteItem], list[tuple]]:
    """生成行/列删除计划：输出到 out_dir，文件名保持原样（重名自动去重）。

    mode/terms 仅用于接口一致性，命名不依赖删除条件。
    """
    used: set[str] = set()
    plan: list[DeleteItem] = []
    errors: list[tuple] = []
    for f in files:
        f = Path(f)
        try:
            dst = build_dest_name(f.stem, f.suffix.lower(), out_dir, used)
            used.add(dst.name.lower())
            plan.append(DeleteItem(src=f, dst=dst))
        except Exception as e:
            errors.append((f, str(e)))
    return plan, errors


def count_delete(src, mode: str, terms: list, case_sensitive: bool) -> tuple:
    """统计要删除的行/列数与总数（供预览）。返回 (删除数, 总数)。"""
    if Path(src).suffix.lower() == '.xls':
        book = xlrd.open_workbook(str(src))
        try:
            sheet = book.sheet_by_index(0)
            dr, dc = _xls_deleted_rows_cols(sheet, mode, terms, case_sensitive)
            if mode == 'row':
                return len(dr), sheet.nrows
            return len(dc), sheet.ncols
        finally:
            book.release_resources()
    wb = load_workbook(str(src), read_only=True)
    try:
        ws = wb.worksheets[0]
        if mode == 'row':
            deleted = _rows_to_delete(ws, terms, case_sensitive)
            return len(deleted), ws.max_row
        deleted = _cols_to_delete(ws, terms, case_sensitive)
        return len(deleted), ws.max_column
    finally:
        wb.close()


def execute_delete(plan: list[DeleteItem], mode: str, terms: list, case_sensitive: bool,
                   log: LogFn, progress: Optional[ProgressFn] = None,
                   cancel_event: Optional[Event] = None) -> dict:
    """执行行/列删除计划（terms = [(值, 匹配方式), ...]）。返回 {'ok': n, 'skip': n, 'fail': n}。"""
    if not plan:
        return {'ok': 0, 'skip': 0, 'fail': 0}
    plan[0].dst.parent.mkdir(parents=True, exist_ok=True)
    total = len(plan)
    ok = skip = fail = 0
    for i, item in enumerate(plan):
        if cancel_event is not None and cancel_event.is_set():
            log('[取消] 用户取消操作')
            break
        try:
            if item.src.suffix.lower() == '.xls':
                _delete_xls(item.src, item.dst, mode, terms, case_sensitive)
            else:
                _delete_openpyxl(item.src, item.dst, mode, terms, case_sensitive)
            ok += 1
            log(f'[OK] {item.src.name} → {item.dst.name}')
        except PermissionError:
            skip += 1
            log(f'[跳过] {item.src.name}：文件被占用，已跳过')
        except Exception as e:
            fail += 1
            log(f'[失败] {item.src.name}：{e}')
        if progress is not None:
            progress(i + 1, total)
    return {'ok': ok, 'skip': skip, 'fail': fail}


# ---------------------------------------------------------------------------
# 按端口号拆分 + 按流量排序 + 标红指定行
# ---------------------------------------------------------------------------

DEFAULT_PORT_HEADER = '集客系统端口号'
DEFAULT_INFLOW_HEADER = '流入速度'
DEFAULT_OUTFLOW_HEADER = '流出速度'
DEFAULT_HIGHLIGHT_ROW = 448

# xlsx 内部 XML 命名空间
_XLSX_MAIN_NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_XLSX_REL_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


@dataclass
class SplitItem:
    src: Path
    dst: Path


class _RowRecord:
    """源表的一行：值、样式 id、流入/流出数值。"""

    __slots__ = ('src_row', 'values', 'styles', 'inflow_num', 'outflow_num')

    def __init__(self, src_row: int, values: tuple, styles: tuple,
                 inflow_num, outflow_num):
        self.src_row = src_row
        self.values = values
        self.styles = styles
        self.inflow_num = inflow_num
        self.outflow_num = outflow_num


@dataclass
class _GroupPlan:
    value: str
    records: list          # 已排序的 _RowRecord
    inflow_total: float
    outflow_total: float
    sort_col: int
    sort_name: str
    invalid_sort_values: int
    red_record: Optional[_RowRecord]


@dataclass
class _Prepared:
    sheet_title: str
    header_values: tuple
    header_styles: tuple
    max_col: int
    total_rows: int
    port_col: int
    inflow_col: int
    outflow_col: int
    highlight_row: int
    groups: list
    style_plans: dict
    layout: dict


@dataclass
class SplitGroupInfo:
    value: str
    row_count: int
    inflow_total: float
    outflow_total: float
    sort_name: str
    invalid_sort_values: int
    red_source_row: Optional[int]


@dataclass
class SplitAnalysis:
    path: Path
    sheet_title: str
    total_rows: int
    highlight_row: int
    groups: list
    summary_rows: int


def _norm_header(value) -> str:
    """表头归一化：去掉所有空白、忽略大小写，便于匹配。"""
    return re.sub(r'\s+', '', value_to_text(value)).casefold()


def _at(seq, one_based_idx: int):
    """取 1 基下标处的值，越界返回 None。"""
    i = one_based_idx - 1
    return seq[i] if 0 <= i < len(seq) else None


def _display_port(value: str) -> str:
    return value if value else '（空）'


def _fmt_number(value: float) -> str:
    """预览/日志中的数字格式：整数不带小数，其余保留两位。"""
    if abs(value - round(value)) < 1e-9:
        return f'{int(round(value)):,}'
    return f'{value:,.2f}'


def _natural_key(text: str) -> list:
    """自然排序键：P2 排在 P10 前面。"""
    parts = re.split(r'(\d+)', text)
    return [(1, p.casefold()) if not p.isdigit() else (0, int(p)) for p in parts]


def _safe_sheet_name(value: str, used: set) -> str:
    """生成合法且不重复的工作表名（Excel 限制 31 字符，禁止 []:*?/\\）。"""
    name = value_to_text(value)
    name = CONTROL_CHARS_RE.sub('', name)
    name = re.sub(r'[\[\]:*?/\\]', '_', name)
    name = name.strip().strip("'").strip()
    if not name:
        name = '空值'
    name = name[:31].rstrip("' ")
    if not name:
        name = '空值'
    base = name
    n = 2
    while name.casefold() in used:
        suffix = f'({n})'
        name = base[:31 - len(suffix)].rstrip("' ") + suffix
        n += 1
    used.add(name.casefold())
    return name


def _capture_style(cell, cache: dict):
    """记录单元格的样式 id，并把样式对象的副本缓存起来（源工作簿关闭后仍可用）。

    ReadOnlyCell 只暴露 _style_id；EmptyCell 没有该属性，按默认样式处理。
    """
    sid = getattr(cell, '_style_id', None)
    if sid is None:
        return None
    if sid not in cache:
        cache[sid] = (
            copy(cell.font), copy(cell.fill), copy(cell.border),
            copy(cell.alignment), copy(cell.protection), cell.number_format,
        )
    return sid


# openpyxl 默认样式（用于判断哪些属性无需逐单元格写入，显著加快大表写出）
_DEFAULT_FONT = Font(name='Calibri', sz=11, family=2,
                     color=Color(theme=1), scheme='minor')
_DEFAULT_FILL = PatternFill()
_DEFAULT_ALIGNMENT = Alignment()
_DEFAULT_PROTECTION = Protection()


def _border_is_empty(border) -> bool:
    if border is None:
        return True
    for side in (border.left, border.right, border.top, border.bottom,
                 border.diagonal, border.vertical, border.horizontal,
                 border.start, border.end):
        if side is not None and side.style is not None:
            return False
    return True


def _style_plan(entry) -> tuple:
    """把样式拆成"需要写入的属性"：与默认值相同的属性返回 None，写出时跳过。

    大表里绝大多数单元格是默认样式，跳过这些赋值可以把写出速度提高数倍。
    """
    font, fill, border, alignment, protection, number_format = entry
    return (
        None if font == _DEFAULT_FONT else font,
        None if fill == _DEFAULT_FILL else fill,
        None if _border_is_empty(border) else border,
        None if alignment == _DEFAULT_ALIGNMENT else alignment,
        None if protection == _DEFAULT_PROTECTION else protection,
        None if number_format in (None, '', 'General') else number_format,
    )


def _read_sheet_layout(path) -> dict:
    """从 xlsx/xlsm 的 XML 头部读取列宽、隐藏列、冻结窗格。

    列宽/窗格在 sheetData 之前，只解压前 512KB 即可，避免大表整段解析。
    读取失败时返回空布局，不影响主流程。
    """
    layout = {'widths': {}, 'hidden': set(), 'freeze_panes': None}
    try:
        with zipfile.ZipFile(str(path)) as z:
            wb_xml = ET.fromstring(z.read('xl/workbook.xml'))
            sheet = wb_xml.find(f'{_XLSX_MAIN_NS}sheets/{_XLSX_MAIN_NS}sheet')
            if sheet is None:
                return layout
            rid = sheet.get(f'{_XLSX_REL_NS}id')
            rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
            target = None
            for rel in rels:
                if rel.get('Id') == rid:
                    target = rel.get('Target')
                    break
            if not target:
                return layout
            target = target.lstrip('/')
            if not target.startswith('xl/'):
                target = 'xl/' + target
            with z.open(target) as f:
                text = f.read(512 * 1024).decode('utf-8', errors='ignore')

        cols = re.search(r'<cols>(.*?)</cols>', text, re.S)
        if cols:
            for col in re.finditer(r'<col\b[^>]*/>', cols.group(1)):
                attrs = dict(re.findall(r'([\w:]+)="([^"]*)"', col.group(0)))
                try:
                    mn = int(attrs.get('min', '1'))
                    mx = int(attrs.get('max', str(mn)))
                except ValueError:
                    continue
                width = attrs.get('width')
                hidden = attrs.get('hidden') in ('1', 'true')
                for i in range(mn, mx + 1):
                    if width is not None:
                        try:
                            layout['widths'][i] = float(width)
                        except ValueError:
                            pass
                    if hidden:
                        layout['hidden'].add(i)

        pane = re.search(r'<pane\b[^>]*/>', text)
        if pane:
            attrs = dict(re.findall(r'([\w:]+)="([^"]*)"', pane.group(0)))
            if attrs.get('state') in ('frozen', 'frozenSplit'):
                top = attrs.get('topLeftCell')
                if top:
                    layout['freeze_panes'] = top
                else:
                    try:
                        x = int(attrs.get('xSplit') or 0)
                        y = int(attrs.get('ySplit') or 0)
                        layout['freeze_panes'] = f'{get_column_letter(x + 1)}{y + 1}'
                    except ValueError:
                        pass
    except Exception:
        return {'widths': {}, 'hidden': set(), 'freeze_panes': None}
    return layout


def _find_header_col(header_norms: list, wanted: str, label: str) -> int:
    target = _norm_header(wanted)
    if not target:
        raise ValueError(f'{label}列名不能为空')
    for i, h in enumerate(header_norms, start=1):
        if h == target:
            return i
    shown = '、'.join(t for t in header_norms if t)[:120]
    raise ValueError(f'未找到{label}列「{wanted}」。当前表头：{shown or "（空）"}')


def _scan_prepare(path, port_header: str, inflow_header: str, outflow_header: str,
                  highlight_row: int) -> _Prepared:
    """读取源文件并完成分组、求和、排序，返回可直接写出的准备结果。"""
    path = Path(path)
    if path.suffix.lower() == '.xls':
        raise ValueError('该功能暂不支持 .xls，请先在 Excel 中另存为 .xlsx 后再处理')
    if path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError('仅支持 .xlsx / .xlsm 文件')
    try:
        highlight_row = int(highlight_row)
    except (TypeError, ValueError):
        raise ValueError('标红行号必须是整数')
    if highlight_row < 2:
        raise ValueError('标红行号必须 ≥ 2（第 1 行是表头）')

    wb = load_workbook(str(path), data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(min_row=1, values_only=False)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            raise ValueError('第一个工作表为空')
        if not header_cells:
            raise ValueError('第一个工作表没有表头')

        style_cache: dict = {}
        header_values = [c.value for c in header_cells]
        header_styles = [_capture_style(c, style_cache) for c in header_cells]
        max_col = max(int(ws.max_column or 0), len(header_values))

        header_norms = [_norm_header(v) for v in header_values]
        port_col = _find_header_col(header_norms, port_header, '端口号')
        inflow_col = _find_header_col(header_norms, inflow_header, '流入速度')
        outflow_col = _find_header_col(header_norms, outflow_header, '流出速度')
        if len({port_col, inflow_col, outflow_col}) < 3:
            raise ValueError('端口号、流入速度、流出速度必须对应三个不同的列')

        groups: dict[str, list] = {}
        inflow_totals: dict[str, float] = {}
        outflow_totals: dict[str, float] = {}
        total_rows = 0

        for src_row, row in enumerate(rows_iter, start=2):
            values = tuple(c.value for c in row)
            if len(values) > max_col:
                max_col = len(values)
            if not any(v is not None and str(v).strip() != '' for v in values):
                continue
            styles = tuple(_capture_style(c, style_cache) for c in row)
            key = value_to_text(_at(values, port_col)).strip()
            inflow_num = _to_number(_at(values, inflow_col))
            outflow_num = _to_number(_at(values, outflow_col))
            groups.setdefault(key, []).append(
                _RowRecord(src_row, values, styles, inflow_num, outflow_num))
            inflow_totals[key] = inflow_totals.get(key, 0.0) + (inflow_num or 0.0)
            outflow_totals[key] = outflow_totals.get(key, 0.0) + (outflow_num or 0.0)
            total_rows += 1
    finally:
        wb.close()

    if total_rows == 0:
        raise ValueError('没有找到任何数据行')

    plans: list = []
    for key in sorted(groups, key=_natural_key):
        records = groups[key]
        inflow_total = inflow_totals.get(key, 0.0)
        outflow_total = outflow_totals.get(key, 0.0)
        # 总和相等时按流入速度排序（与"取较大者"的常规理解一致）
        if inflow_total >= outflow_total:
            sort_col, sort_name = inflow_col, (value_to_text(_at(header_values, inflow_col)).strip()
                                               or inflow_header)
        else:
            sort_col, sort_name = outflow_col, (value_to_text(_at(header_values, outflow_col)).strip()
                                                or outflow_header)

        def sort_key(rec, _col=sort_col):
            n = rec.inflow_num if _col == inflow_col else rec.outflow_num
            # 数值在前按降序，非数值/空值排在最后（sort 稳定，保持原顺序）
            return (n is None, -(n if n is not None else 0.0))

        records.sort(key=sort_key)
        invalid = sum(1 for rec in records
                      if (rec.inflow_num if sort_col == inflow_col else rec.outflow_num) is None)
        red_record = records[highlight_row - 2] if len(records) >= highlight_row - 1 else None
        plans.append(_GroupPlan(
            value=key,
            records=records,
            inflow_total=inflow_total,
            outflow_total=outflow_total,
            sort_col=sort_col,
            sort_name=sort_name,
            invalid_sort_values=invalid,
            red_record=red_record,
        ))

    return _Prepared(
        sheet_title=ws.title,
        header_values=tuple(header_values),
        header_styles=tuple(header_styles),
        max_col=max_col,
        total_rows=total_rows,
        port_col=port_col,
        inflow_col=inflow_col,
        outflow_col=outflow_col,
        highlight_row=highlight_row,
        groups=plans,
        style_plans={sid: _style_plan(entry) for sid, entry in style_cache.items()},
        layout=_read_sheet_layout(path),
    )


def analyze_split(path, port_header: str = DEFAULT_PORT_HEADER,
                  inflow_header: str = DEFAULT_INFLOW_HEADER,
                  outflow_header: str = DEFAULT_OUTFLOW_HEADER,
                  highlight_row: int = DEFAULT_HIGHLIGHT_ROW) -> SplitAnalysis:
    """分析拆分结果（供预览），不写出文件。"""
    prepared = _scan_prepare(path, port_header, inflow_header, outflow_header, highlight_row)
    infos = [
        SplitGroupInfo(
            value=g.value,
            row_count=len(g.records),
            inflow_total=g.inflow_total,
            outflow_total=g.outflow_total,
            sort_name=g.sort_name,
            invalid_sort_values=g.invalid_sort_values,
            red_source_row=g.red_record.src_row if g.red_record is not None else None,
        )
        for g in prepared.groups
    ]
    return SplitAnalysis(
        path=Path(path),
        sheet_title=prepared.sheet_title,
        total_rows=prepared.total_rows,
        highlight_row=prepared.highlight_row,
        groups=infos,
        summary_rows=sum(1 for g in infos if g.red_source_row is not None),
    )


def preview_split(path, port_header: str = DEFAULT_PORT_HEADER,
                  inflow_header: str = DEFAULT_INFLOW_HEADER,
                  outflow_header: str = DEFAULT_OUTFLOW_HEADER,
                  highlight_row: int = DEFAULT_HIGHLIGHT_ROW) -> list[str]:
    """生成预览文本行。"""
    a = analyze_split(path, port_header, inflow_header, outflow_header, highlight_row)
    lines = [f'{a.path.name}：{a.total_rows} 行数据，{len(a.groups)} 个端口号']
    for g in a.groups:
        lines.append(
            f'  · {_display_port(g.value)}：{g.row_count} 行；'
            f'流入总和 {_fmt_number(g.inflow_total)}，流出总和 {_fmt_number(g.outflow_total)}'
            f' → 按「{g.sort_name}」降序')
        if g.red_source_row is None:
            lines.append(f'    第 {a.highlight_row} 行：数据不足 {a.highlight_row - 1} 行，不标红')
        else:
            lines.append(f'    第 {a.highlight_row} 行：原第 {g.red_source_row} 行 → 标红并汇总')
        if g.invalid_sort_values:
            lines.append(f'    提示：排序列有 {g.invalid_sort_values} 个非数值，已排在最后')
    lines.append(
        f'  → 输出 {len(a.groups)} 个子表 + 1 个「{a.highlight_row}行汇总」，'
        f'汇总标红行 {a.summary_rows} 行')
    return lines


def plan_split(files: Iterable, out_dir,
               suffix: str = '_按端口拆分') -> tuple[list[SplitItem], list[tuple]]:
    """生成拆分计划：每个源文件输出一个 .xlsx（内含多个子表 + 汇总表）。"""
    used: set[str] = set()
    plan: list[SplitItem] = []
    errors: list[tuple] = []
    for f in files:
        f = Path(f)
        if f.suffix.lower() not in ('.xlsx', '.xlsm'):
            errors.append((f, '仅支持 .xlsx / .xlsm 文件'))
            continue
        try:
            dst = build_dest_name(f'{f.stem}{suffix}', '.xlsx', out_dir, used)
            used.add(dst.name.lower())
            plan.append(SplitItem(src=f, dst=dst))
        except Exception as e:
            errors.append((f, str(e)))
    return plan, errors


def _apply_layout(ws, layout: dict, max_col: int):
    """把列宽、隐藏列、冻结窗格应用到输出工作表。"""
    if not layout:
        return
    widths = layout.get('widths') or {}
    hidden = layout.get('hidden') or set()
    for i in range(1, max_col + 1):
        letter = get_column_letter(i)
        if i in widths:
            ws.column_dimensions[letter].width = widths[i]
        if i in hidden:
            ws.column_dimensions[letter].hidden = True
    if layout.get('freeze_panes'):
        ws.freeze_panes = layout['freeze_panes']


def _make_row(ws, values, styles, prepared: _Prepared, red_font_cache=None):
    """构造一行 WriteOnlyCell；red_font_cache 非 None 时整行字体标红。"""
    cells = []
    for i in range(prepared.max_col):
        cell = WriteOnlyCell(ws, value=_at(values, i + 1))
        sid = _at(styles, i + 1)
        plan = prepared.style_plans.get(sid)
        if plan is not None:
            font, fill, border, alignment, protection, number_format = plan
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if alignment is not None:
                cell.alignment = alignment
            if protection is not None:
                cell.protection = protection
            if number_format is not None:
                cell.number_format = number_format
        if red_font_cache is not None:
            red_font = red_font_cache.get(sid)
            if red_font is None:
                red_font = copy(cell.font)
                red_font.color = Color(rgb='FFFF0000')
                red_font_cache[sid] = red_font
            cell.font = red_font
        cells.append(cell)
    return cells


def _write_output(prepared: _Prepared, dst, log: LogFn):
    """把准备结果写成一个 .xlsx：每个端口号一个子表 + 一个标红行汇总表。"""
    out = Workbook(write_only=True)
    used_titles: set[str] = set()
    red_font_cache: dict = {}

    for g in prepared.groups:
        title = _safe_sheet_name(g.value, used_titles)
        ws = out.create_sheet(title)
        _apply_layout(ws, prepared.layout, prepared.max_col)
        ws.append(_make_row(ws, prepared.header_values, prepared.header_styles, prepared))
        for rec in g.records:
            is_red = rec is g.red_record
            ws.append(_make_row(ws, rec.values, rec.styles, prepared,
                                red_font_cache if is_red else None))
        if g.red_record is None:
            note = f'数据不足 {prepared.highlight_row - 1} 行，未标红'
        else:
            note = f'第 {prepared.highlight_row} 行已标红'
        log(f'[信息] 子表「{title}」：{len(g.records)} 行；'
            f'流入总和 {_fmt_number(g.inflow_total)}，流出总和 {_fmt_number(g.outflow_total)}；'
            f'按「{g.sort_name}」降序；{note}')

    summary_title = _safe_sheet_name(f'{prepared.highlight_row}行汇总', used_titles)
    sws = out.create_sheet(summary_title)
    _apply_layout(sws, prepared.layout, prepared.max_col)
    sws.append(_make_row(sws, prepared.header_values, prepared.header_styles, prepared))
    for g in prepared.groups:
        if g.red_record is not None:
            sws.append(_make_row(sws, g.red_record.values, g.red_record.styles,
                                 prepared, red_font_cache))
    log(f'[信息] 汇总表「{summary_title}」：{sum(1 for g in prepared.groups if g.red_record)} 行')

    out.save(str(dst))


def execute_split(plan: list[SplitItem], port_header: str, inflow_header: str,
                  outflow_header: str, highlight_row: int, log: LogFn,
                  progress: Optional[ProgressFn] = None,
                  cancel_event: Optional[Event] = None) -> dict:
    """执行拆分计划。返回 {'ok': n, 'skip': n, 'fail': n}。"""
    if not plan:
        return {'ok': 0, 'skip': 0, 'fail': 0}
    plan[0].dst.parent.mkdir(parents=True, exist_ok=True)
    total = len(plan)
    ok = skip = fail = 0
    for i, item in enumerate(plan):
        if cancel_event is not None and cancel_event.is_set():
            log('[取消] 用户取消操作')
            break
        try:
            log(f'[信息] 正在读取并分析 {item.src.name} …')
            prepared = _scan_prepare(item.src, port_header, inflow_header,
                                     outflow_header, highlight_row)
            log(f'[信息] 共 {prepared.total_rows} 行、{len(prepared.groups)} 个端口号，开始写出 …')
            _write_output(prepared, item.dst, log)
            ok += 1
            log(f'[OK] {item.src.name} → {item.dst.name}')
        except PermissionError:
            skip += 1
            log(f'[跳过] {item.src.name}：文件被占用，已跳过')
        except Exception as e:
            fail += 1
            log(f'[失败] {item.src.name}：{e}')
        if progress is not None:
            progress(i + 1, total)
    return {'ok': ok, 'skip': skip, 'fail': fail}
