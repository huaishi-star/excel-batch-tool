"""GUI 全流程自动化测试（开发用）。

模拟真实操作：添加文件 → 设置规则 → 预览 → 执行 → 校验结果与日志。
后台线程管道通过手动 pump tab._poll() 驱动（无需 mainloop）。

用法：python test_gui_flow.py
"""

import shutil
import tempfile
import time
from pathlib import Path

import app as appmod
import config_store

BASE = Path(__file__).parent
SAMPLES = BASE / 'test_samples'
OUT = BASE / 'test_out_gui'


def pump_until_idle(tab, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        time.sleep(0.05)
        tab._poll()
        if tab._worker is None:
            return True
    return False


def run_rename(app):
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir()
    tab = app.rename_tab
    tab.picker.files = [str(SAMPLES / f) for f in ('数据1.xlsx', '数据2.xlsx')]
    tab.picker._refresh()
    for c in ('A1', 'B2', 'C3'):
        tab.coord_entry.insert(0, c)
        tab._add_coord()
    tab.sep_entry.insert(0, '-')
    tab.out_var.set(str(OUT))

    # 预览
    tab._do_preview()
    assert pump_until_idle(tab), '重命名预览超时'
    preview = tab.preview_text.get('1.0', 'end')
    assert '张三-20240115-98.5.xlsx' in preview, preview
    assert '李四-20240220-100.xlsx' in preview, preview
    print('  [预览] 两个输出名正确')

    # 执行
    tab._run()
    assert pump_until_idle(tab), '重命名执行超时'
    assert (OUT / '张三-20240115-98.5.xlsx').exists(), '输出文件1不存在'
    assert (OUT / '李四-20240220-100.xlsx').exists(), '输出文件2不存在'
    log = tab.log_text.get('1.0', 'end')
    assert log.count('[OK]') == 2 and '成功 2 个' in log, log
    assert tab.run_button.instate(['!disabled']), '完成后按钮应恢复可用'
    print('  [执行] 两个文件复制成功，按钮恢复，日志正确')
    print('重命名 GUI 流程 PASS')


def run_fill(app):
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir()
    tab = app.fill_tab
    tab.picker.files = [str(SAMPLES / f) for f in ('数据1.xlsx', '数据2.xlsx')]
    tab.picker._refresh()
    tab.template_var.set(str(SAMPLES / '合同模板.xlsx'))
    for s, t in (('A1', 'B2'), ('B2', 'C5'), ('C3', 'D6')):
        tab.src_entry.insert(0, s)
        tab.tpl_entry.insert(0, t)
        tab._add_map()
    tab.out_var.set(str(OUT))

    # 预览
    tab._do_preview()
    assert pump_until_idle(tab), '填充预览超时'
    preview = tab.preview_text.get('1.0', 'end')
    assert '合同模板_数据1.xlsx' in preview and '合同模板_数据2.xlsx' in preview, preview
    print('  [预览] 输出名正确')

    # 执行
    tab._run()
    assert pump_until_idle(tab), '填充执行超时'
    assert (OUT / '合同模板_数据1.xlsx').exists(), '填充输出1不存在'
    assert (OUT / '合同模板_数据2.xlsx').exists(), '填充输出2不存在'
    log = tab.log_text.get('1.0', 'end')
    assert log.count('[OK]') == 2 and '成功 2 个' in log, log

    # 校验内容
    from datetime import datetime
    from openpyxl import load_workbook
    wb = load_workbook(OUT / '合同模板_数据1.xlsx')
    ws = wb.worksheets[0]
    assert ws['B2'].value == '张三', ws['B2'].value
    assert ws['C5'].value == datetime(2024, 1, 15), ws['C5'].value
    assert ws['B2'].number_format == '0.00', ws['B2'].number_format
    wb.close()
    print('  [执行] 两个文件生成，映射值与格式正确')
    print('模板填充 GUI 流程 PASS')


def run_fill_column(app):
    """整列映射 GUI 流程：源列数据填入模板列，已有内容被推下不覆盖。"""
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir()
    tab = app.fill_tab
    tab.picker.files = [str(SAMPLES / '列数据源.xlsx')]
    tab.picker._refresh()
    tab.template_var.set(str(SAMPLES / '表格模板.xlsx'))
    tab.map_type.set('整列 → 整列')
    tab._toggle_map_type()
    tab.src_col_entry.insert(0, 'B')
    tab.src_row_entry.insert(0, '2')
    tab.tpl_col_entry.insert(0, 'A')
    tab.tpl_row_entry.insert(0, '6')
    tab._add_map()
    tab.out_var.set(str(OUT))
    tab._run()
    assert pump_until_idle(tab), '整列填充执行超时'

    out = OUT / '表格模板_列数据源.xlsx'
    assert out.exists(), '整列填充输出不存在'
    from openpyxl import load_workbook
    ws = load_workbook(out).worksheets[0]
    assert ws['A6'].value == '产品A', ws['A6'].value
    assert ws['A7'].value == '产品B', ws['A7'].value
    assert ws['A10'].value == '产品E', ws['A10'].value
    assert ws['A11'].value == '旧项1', ws['A11'].value   # 被推下，未被覆盖
    assert ws['A12'].value == '旧项2', ws['A12'].value
    log = tab.log_text.get('1.0', 'end')
    assert '[OK]' in log and '成功 1 个' in log, log
    print('  [整列] 源列填入模板列，已有内容下移不覆盖')
    print('模板填充-整列 GUI 流程 PASS')


def run_delete(app):
    """行/列删除 GUI 流程：预览统计 + 执行删除。"""
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir()
    # 造一个源文件
    from openpyxl import Workbook
    src = OUT / '待删.xlsx'
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '项目'
    ws['A2'] = '作废单'
    ws['A3'] = '正常单'
    ws['A4'] = '作废2'
    wb.save(src)

    tab = app.delete_tab
    tab.picker.files = [str(src)]
    tab.picker._refresh()
    tab.target_entry.insert(0, '作废')
    tab._add_term()  # 添加条件"作废(包含)"
    tab.out_var.set(str(OUT))

    # 预览
    tab._do_preview()
    assert pump_until_idle(tab), '删除预览超时'
    preview = tab.preview_text.get('1.0', 'end')
    assert '将删除 2 行（共 4 行）' in preview, preview
    print('  [预览] 统计 2 行正确')

    # 执行（输出与源同目录，会触发改名提示；弹窗需放行）
    appmod.messagebox.askyesno = lambda *a, **k: True
    tab._run()
    assert pump_until_idle(tab), '删除执行超时'
    out_files = [p.name for p in OUT.glob('*.xlsx') if p.name != '待删.xlsx']
    assert out_files, '删除输出文件不存在'
    from openpyxl import load_workbook
    ws2 = load_workbook(OUT / out_files[0]).worksheets[0]
    col = [ws2.cell(row=r, column=1).value for r in range(1, ws2.max_row + 1)]
    assert col == ['项目', '正常单'], col
    print('  [执行] 作废行已删除')
    print('行/列删除 GUI 流程 PASS')


def run_split(app):
    """端口拆分 GUI 流程：预览排序依据 + 执行拆分/标红/汇总。"""
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir()
    from openpyxl import Workbook, load_workbook
    src = OUT / '拆分源.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['集客系统端口号', '流入速度', '流出速度', '备注'])
    for row in [('P001', 10, 40, 'a'), ('P001', 20, 100, 'b'), ('P001', 30, 20, 'c'),
                ('P001', 40, 80, 'd'), ('P001', 50, 0, 'e'), ('P001', 60, 60, 'f'),
                ('P002', 5, 700, 'g'), ('P002', 15, 1000, 'h')]:
        ws.append(row)
    wb.save(src)

    tab = app.split_tab
    tab.picker.files = [str(src)]
    tab.picker._refresh()
    tab.highlight_var.set('5')
    tab.out_var.set(str(OUT))

    tab._do_preview()
    assert pump_until_idle(tab), '端口拆分预览超时'
    preview = tab.preview_text.get('1.0', 'end')
    assert 'P001' in preview and '流出速度' in preview and '标红并汇总' in preview, preview
    print('  [预览] 端口分组与排序依据正确')

    tab._run()
    assert pump_until_idle(tab), '端口拆分执行超时'
    out = OUT / '拆分源_按端口拆分.xlsx'
    assert out.exists(), '端口拆分输出不存在'
    o = load_workbook(out)
    assert o.sheetnames == ['P001', 'P002', '5行汇总'], o.sheetnames
    assert [o['P001'].cell(row=5, column=c).value for c in range(1, 5)] == ['P001', 10, 40, 'a'], \
        '第5行数据异常'
    assert all(o['P001'].cell(row=5, column=c).font.color.rgb == 'FFFF0000'
               for c in range(1, 5)), '第5行未标红'
    assert o['5行汇总'].max_row == 2, '汇总表行数异常（P002 数据不足，不汇总）'
    o.close()
    log = tab.log_text.get('1.0', 'end')
    assert '[OK]' in log and '成功 1 个' in log, log
    print('  [执行] 2 个子表 + 1 个汇总表，第5行标红')
    print('端口拆分 GUI 流程 PASS')


def test_profiles():
    """规则/映射档案：保存 → 载入 → 删除；不保存文件列表/输出文件夹。"""
    with tempfile.TemporaryDirectory() as tmp:
        config_store.set_base_dir(tmp)
        appmod.simpledialog.askstring = lambda *a, **k: '测试记录'
        try:
            a = appmod.App()

            # 保存重命名规则（单元格 + 文字 + 原文件名混排）
            for c in ('A1', 'B2', 'C3'):
                a.rename_tab.coord_entry.insert(0, c)
                a.rename_tab._add_coord()
            a.rename_tab.text_entry.insert(0, '客户_')
            a.rename_tab._add_text()
            a.rename_tab._add_filename()
            a.rename_tab.sep_entry.insert(0, '-')
            a.rename_tab._save_rule()
            assert config_store.list_profiles('rename') == ['测试记录'], config_store.list_profiles('rename')
            prof = config_store.get_profile('rename', '测试记录')
            assert prof['rule'] == [['cell', 'A1'], ['cell', 'B2'], ['cell', 'C3'],
                                    ['text', '客户_'], ['filename', '']], prof
            assert prof['separator'] == '-', prof
            assert 'files' not in prof and 'out_dir' not in prof, '不应保存文件列表/输出文件夹'
            print('  [规则] 保存正确（含文字/原文件名），不含文件列表/输出文件夹')

            # 载入重命名规则（先清空再载入）
            a.rename_tab._rule_items.clear()
            a.rename_tab._refresh_rule_list()
            a.rename_tab.sep_entry.delete(0, 'end')
            a.rename_tab.rule_cb.set('测试记录')
            a.rename_tab._load_rule()
            assert a.rename_tab._rule_items == [['cell', 'A1'], ['cell', 'B2'], ['cell', 'C3'],
                                                ['text', '客户_'], ['filename', '']], a.rename_tab._rule_items
            assert a.rename_tab.sep_entry.get() == '-'
            print('  [规则] 载入正确（含文字/原文件名）')

            # 保存映射（单元格 + 整列混排）
            a.fill_tab.src_entry.insert(0, 'A1')
            a.fill_tab.tpl_entry.insert(0, 'B2')
            a.fill_tab._add_map()
            a.fill_tab.map_type.set('整列 → 整列')
            a.fill_tab._toggle_map_type()
            a.fill_tab.src_col_entry.insert(0, 'C')
            a.fill_tab.src_row_entry.insert(0, '2')
            a.fill_tab.tpl_col_entry.insert(0, 'B')
            a.fill_tab.tpl_row_entry.insert(0, '6')
            a.fill_tab._add_map()
            a.fill_tab._save_mapping()
            mprof = config_store.get_profile('fill', '测试记录')
            assert mprof['mapping'] == [['cell', 'A1', 'B2'], ['column', 'C', 2, 'B', 6]], mprof
            assert 'files' not in mprof and 'out_dir' not in mprof, '不应保存文件列表/输出文件夹'
            print('  [映射] 保存正确（含整列映射），不含文件列表/输出文件夹')

            # 载入映射
            a.fill_tab._mappings.clear()
            a.fill_tab._refresh_map_tree()
            a.fill_tab.map_cb.set('测试记录')
            a.fill_tab._load_mapping()
            assert a.fill_tab._mappings == [['cell', 'A1', 'B2'], ['column', 'C', 2, 'B', 6]], \
                a.fill_tab._mappings
            print('  [映射] 载入正确（含整列映射）')

            # 删除规则保存/载入
            a.delete_tab.target_entry.insert(0, '作废')
            a.delete_tab._add_term()
            a.delete_tab.target_entry.insert(0, '合计')
            a.delete_tab.match_combo.set('等于')
            a.delete_tab._add_term()
            a.delete_tab.mode_var.set('column')
            a.delete_tab._save_rule()
            dprof = config_store.get_profile('delete', '测试记录')
            assert dprof['mode'] == 'column', dprof
            assert dprof['terms'] == [['作废', 'contains'], ['合计', 'equals']], dprof
            assert dprof['case_sensitive'] is False, dprof
            a.delete_tab._terms.clear()
            a.delete_tab._refresh_terms()
            a.delete_tab.rule_cb.set('测试记录')
            a.delete_tab._load_rule()
            assert a.delete_tab._terms == [('作废', 'contains'), ('合计', 'equals')], a.delete_tab._terms
            assert a.delete_tab.mode_var.get() == 'column', a.delete_tab.mode_var.get()
            print('  [删除规则] 保存/载入正确')

            # 删除（确认弹窗返回 True）
            appmod.messagebox.askyesno = lambda *a, **k: True
            a.rename_tab.rule_cb.set('测试记录')
            a.rename_tab._delete_rule()
            assert '测试记录' not in config_store.list_profiles('rename'), config_store.list_profiles('rename')
            print('  [删除] 正确')
            a.destroy()
        finally:
            config_store.set_base_dir(None)
    print('规则/映射 档案测试 PASS')


def main():
    a = appmod.App()
    try:
        run_rename(a)
        run_fill(a)
        run_fill_column(a)
        run_delete(a)
        run_split(a)
        test_profiles()
        print('\n全部 GUI 流程测试通过 ✅')
    finally:
        a.destroy()


if __name__ == '__main__':
    main()
