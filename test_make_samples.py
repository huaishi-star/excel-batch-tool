"""生成测试样例文件到 test_samples\（仅开发调试用，不随软件分发）。

用法：python test_make_samples.py

覆盖的重命名测试点：
- 数据1：正常值 → 张三-20240115-98.5.xlsx
- 数据2 / 数据3：内容相同 → 触发重名去重 (1)
- 数据4：含 Windows 非法字符 → 替换为下划线
- 数据5：超长文本（160 字符）→ 截断到 150
- 数据6：中间单元格为空 → 跳过空值拼接
- 数据7：A1 为 Windows 保留名 CON → 加下划线前缀

注意：宏文件(.xlsm) 无法由本脚本伪造，如需测试宏保留，请用 Excel 手动另存一个。
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

OUT = Path(__file__).parent / 'test_samples'


def make_data(name: str, vals: dict):
    wb = Workbook()
    ws = wb.active
    for coord, v in vals.items():
        ws[coord] = v
    wb.save(OUT / name)


def make_template():
    wb = Workbook()
    ws = wb.active
    # 标题（合并单元格 + 加粗居中）
    ws['A1'] = '合同信息表'
    ws.merge_cells('A1:D1')
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    # 留空的填写区，带边框和数字格式
    ws['A2'] = '客户名称'
    ws['B2'] = None
    ws['B2'].number_format = '0.00'          # 若写入数字，按两位小数显示
    ws['A5'] = '日期'
    ws['C5'] = None
    ws['C5'].number_format = 'yyyy-mm-dd'    # 若写入日期，按该格式显示
    ws['D6'] = None

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for coord in ('B2', 'C5', 'D6'):
        ws[coord].border = border
    wb.save(OUT / '合同模板.xlsx')


def make_column_samples():
    """整列填充测试用：源文件一列数据 + 模板已有内容的表格。"""
    # 列数据源：B 列 2..6 行有数据，第 7 行为空
    wb = Workbook()
    ws = wb.active
    ws['B1'] = '产品'
    for r, v in enumerate(['产品A', '产品B', '产品C', '产品D', '产品E'], start=2):
        ws[f'B{r}'] = v
    wb.save(OUT / '列数据源.xlsx')

    # 表格模板：A 列 6..7 行已有内容（会被推下，不覆盖），第 5 行为表头
    wb = Workbook()
    ws = wb.active
    ws['A5'] = '产品'
    ws['A6'] = '旧项1'
    ws['A7'] = '旧项2'
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for coord in ('A5', 'A6', 'A7'):
        ws[coord].border = border
    wb.save(OUT / '表格模板.xlsx')


def make_xls_samples():
    """旧版 .xls 测试用：数据源 + 带已有内容的表格模板。"""
    import xlwt
    from datetime import datetime
    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sheet1')
    ws.write(0, 0, '张三')
    ws.write(1, 1, datetime(2024, 1, 15), xlwt.easyxf('', num_format_str='YYYY-MM-DD'))
    ws.write(2, 2, 98.5)
    for r, v in enumerate(['产品A', '产品B', '产品C'], start=4):
        ws.write(r, 1, v)   # B 列 5..7 行
    wb.save(OUT / '旧数据.xls')

    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sheet1')
    border = xlwt.easyxf('borders: left thin, right thin, top thin, bottom thin')
    for r, v in zip((4, 5, 6), ('产品', '旧项1', '旧项2')):
        ws.write(r, 0, v, border)   # A5/A6/A7
    wb.save(OUT / '旧表格模板.xls')


def make_split_sample():
    """端口拆分测试用：3 个端口，流入/流出总和大的一方决定排序。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '数据'
    for c, v in enumerate(['集客系统端口号', '流入速度', '流出速度', '备注'], start=1):
        ws.cell(row=1, column=c, value=v)
    ws['A1'].font = Font(bold=True)
    rows = [
        ('P001', 10, 40, 'a'), ('P001', 20, 100, 'b'), ('P001', 30, 20, 'c'),
        ('P001', 40, 80, 'd'), ('P001', 50, 0, 'e'), ('P001', 60, 60, 'f'),
        ('P002', 5, 700, 'g'), ('P002', 15, 1000, 'h'), ('P002', 25, 800, 'i'),
        ('P002', 35, 900, 'j'),
        ('P003', 100, 1, 'k'), ('P003', 'N/A', 3, 'm'), ('P003', 200, 2, 'l'),
    ]
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
    ws.column_dimensions['A'].width = 18
    ws.freeze_panes = 'A2'
    wb.save(OUT / '拆分样例.xlsx')


def main():
    OUT.mkdir(exist_ok=True)
    make_data('数据1.xlsx', {
        'A1': '张三',
        'B2': datetime(2024, 1, 15),
        'C3': 98.5,
        'D4': '<A*B:?',
    })
    make_data('数据2.xlsx', {
        'A1': '李四',
        'B2': datetime(2024, 2, 20),
        'C3': 100.0,
    })
    make_data('数据3.xlsx', {   # 与数据2 完全相同，用于测重名
        'A1': '李四',
        'B2': datetime(2024, 2, 20),
        'C3': 100.0,
    })
    make_data('数据4.xlsx', {   # 非法字符
        'A1': '王五:Q?A',
        'B2': datetime(2024, 3, 1),
        'C3': 7,
    })
    make_data('数据5.xlsx', {   # 超长文本
        'A1': '很' * 160,
        'B2': datetime(2024, 4, 1),
        'C3': 1,
    })
    make_data('数据6.xlsx', {   # 中间单元格为空
        'A1': '赵六',
        'B2': None,
        'C3': 7,
    })
    make_data('数据7.xlsx', {   # 保留设备名
        'A1': 'CON',
        'B2': datetime(2024, 5, 1),
        'C3': 9,
    })
    make_template()
    make_column_samples()
    make_xls_samples()
    make_split_sample()
    print(f'样例文件已生成到：{OUT}')


if __name__ == '__main__':
    main()
